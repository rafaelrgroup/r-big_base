from __future__ import annotations
import os
import re
import asyncio
import json
import secrets
import sqlite3
import csv
import io
import zipfile
from pathlib import Path
from datetime import datetime, timezone, timedelta
from contextlib import asynccontextmanager
from concurrent.futures import ThreadPoolExecutor
from fastapi import FastAPI, Request, Response, HTTPException, UploadFile
from fastapi.responses import JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from starlette.concurrency import run_in_threadpool
from .rate_limit import RedisLimiter, RateLimitUnavailable
from pydantic import BaseModel, ConfigDict, Field
from .store import Store
from .security import Security, Limiter, hashed, PERMISSIONS
from .domain import Enrichment, ItemInput, now, uid, fingerprint, project, timestamp, record_flags
from .enrichment import enrich_entity
from .sorting import resolve_sort, sort_entities, sorting_catalog
from .search import validate_filter, matches, FIELDS, OPS
from .exports import Exporter, prepare_job
from .imports import Importer
from .relationships import connections, resolve as resolve_relationship
from .catalogs import FieldDefinitionInput, create_definition, update_definition, public_definition, bind_definition, VALUE_POLICY
from .ingress_json import IngressJSONError, is_json_content_type, load_preserving_json

class Login(BaseModel):
    username:str=Field(max_length=160);password:str=Field(max_length=256)
class OTP(BaseModel):
    challenge:str=Field(max_length=100);code:str=Field(pattern=r'^\d{6}$')
class StepUpOTP(BaseModel):
    model_config=ConfigDict(extra='forbid')
    code:str=Field(pattern=r'^[0-9]{6}$')


def create_app(root=None,testing=False,*,canonical_reads=None):
    if not testing and os.environ.get('BIGBASE_ENV','development')!='development':
        raise RuntimeError('Este adaptador é de desenvolvimento; configure o backend definitivo antes de produção.')
    root=Path(root or os.environ.get('BIGBASE_DATA','var')).resolve();root.mkdir(parents=True,exist_ok=True);root.chmod(0o700)
    store=Store(root/'development.sqlite3');security=Security(store,root);limiter=RedisLimiter(os.environ['BIGBASE_REDIS_URL']) if os.environ.get('BIGBASE_REDIS_URL') else Limiter();exports=Exporter(store,root/'exports')
    imports=Importer(store)
    executor=ThreadPoolExecutor(max_workers=1,thread_name_prefix='bigbase-jobs')
    @asynccontextmanager
    async def lifespan(app):
        with store.transaction() as c: pending=[j['id'] for j in store.all(c,'job') if j['status'] in {'pending','preparing'}]
        for id in pending:executor.submit(exports.run,id)
        with store.transaction() as c: pending_imports=[j['id'] for j in store.all(c,'import_job') if j['status'] in {'pending','processing'}]
        for id in pending_imports:executor.submit(imports.run,id)
        async def housekeeping():
            while True:
                await run_in_threadpool(exports.cleanup_expired)
                await asyncio.sleep(3600)
        cleaner=asyncio.create_task(housekeeping())
        async def rotation_housekeeping():
            while True:
                await run_in_threadpool(security.cleanup_rotation_receipts)
                await asyncio.sleep(60)
        receipt_cleaner=asyncio.create_task(rotation_housekeeping())
        maintenance_tasks=[cleaner,receipt_cleaner]
        if canonical_reads is not None:
            maintenance_tasks.append(asyncio.create_task(maintain_canonical_cursors(canonical_reads),name='canonical-cursor-cleanup'))
        try:yield
        finally:
            for task in maintenance_tasks:task.cancel()
            for task in maintenance_tasks:
                try:await task
                except asyncio.CancelledError:pass
            executor.shutdown(wait=True)
            if isinstance(limiter,RedisLimiter):limiter.close()
    app=FastAPI(title='BIG BASE',version='0.1.0',docs_url='/api/docs',openapi_url='/api/openapi.json',lifespan=lifespan)
    app.state.store=store;app.state.security=security;app.state.exports=exports;app.state.imports=imports;app.state.job_executor=executor
    with store.transaction() as c:
        if not store.get(c,'source','manual'):store.put(c,'source',{'id':'manual','name':'Cadastro manual','active':True})
    @app.middleware('http')
    async def guard(request,call_next):
        rid=uid();ip=request.client.host if request.client else 'unknown'
        try:
            ip_allowed=testing or await run_in_threadpool(limiter.consume,'ip:'+ip)
            auth_allowed=testing or not (request.url.path.startswith('/api/v1/auth/') and request.method=='POST') or await run_in_threadpool(limiter.consume,'auth:'+ip,rate=1/12,burst=5)
        except RateLimitUnavailable:return JSONResponse({'detail':'Controle de acesso temporariamente indisponível','request_id':rid},503,headers={'Retry-After':'5'})
        if not ip_allowed:return JSONResponse({'detail':'Limite de solicitações','request_id':rid},429,headers={'Retry-After':'1'})
        if not auth_allowed:return JSONResponse({'detail':'Muitas tentativas; aguarde','request_id':rid},429,headers={'Retry-After':'60'})
        if request.method in {'POST','PATCH','PUT'}:
            size=0
            async for chunk in request.stream():
                size+=len(chunk)
                if size>2*1024*1024:return JSONResponse({'detail':'Corpo excede 2 MiB'},413)
                # Keep bounded input available to FastAPI after inspecting actual bytes.
                request._body=getattr(request,'_body',b'')+chunk
        if request.method in {'POST','PATCH','PUT'} and is_json_content_type(request.headers.get('content-type')):
            try:
                if (request.method == 'POST' and request.url.path in {'/api/v1/canonical/people/enrich', '/api/v1/canonical/companies/enrich'}) or (request.method == 'PATCH' and re.fullmatch(r'/api/v1/canonical/(people|companies)/[^/]+/items/[^/]+/value', request.url.path)):
                    from .canonical_store import decode
                    decode(getattr(request,'_body',b''))
                else:
                    load_preserving_json(getattr(request,'_body',b''))
            except IngressJSONError as exc:return JSONResponse({'detail':str(exc),'code':exc.code,'request_id':rid},422)
            except (ValueError,UnicodeDecodeError,RecursionError):return JSONResponse({'detail':'JSON inválido, profundo demais ou com número não finito','request_id':rid},422)
        response=await call_next(request)
        response.headers.update({'X-Request-ID':rid,'X-Content-Type-Options':'nosniff','Referrer-Policy':'same-origin','Cache-Control':'no-store','Content-Security-Policy':"default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; script-src 'self'; connect-src 'self'; frame-ancestors 'none'"})
        return response
    def auth(c,request,scope):
        u,key=security.authorize(c,request,scope)
        if not testing:
            try:
                if not limiter.consume('user:'+u['id'],rate=20,burst=40):raise HTTPException(429,'Limite agregado do usuário',headers={'Retry-After':'1'})
                category='read' if scope=='read' else 'write'
                if not limiter.consume('user:'+u['id']+':'+category,rate=20 if category=='read' else 5,burst=40 if category=='read' else 10):raise HTTPException(429,'Limite da operação',headers={'Retry-After':'1'})
                if key and not limiter.consume('key:'+key['public_id']+':'+category,rate=20 if category=='read' else 5,burst=40 if category=='read' else 10):raise HTTPException(429,'Limite da integração',headers={'Retry-After':'1'})
            except RateLimitUnavailable:raise HTTPException(503,'Coordenador de limites indisponível',headers={'Retry-After':'5'})
        return u,key
    def account_limit(identity):
        if testing:return
        try:
            if not limiter.consume('account:'+hashed(identity),rate=1/12,burst=5):raise HTTPException(429,'Muitas tentativas nesta conta; aguarde',headers={'Retry-After':'60'})
        except RateLimitUnavailable:raise HTTPException(503,'Coordenador de limites indisponível',headers={'Retry-After':'5'})
    def challenge_limit(c,token):
        ch=store.get(c,'challenge',hashed(token))
        if ch:account_limit('otp:'+ch['user_id'])
    def source_check(c,id,key):
        src=store.get(c,'source',id)
        if not src or not src['active']:raise HTTPException(422,'Origem inexistente ou inativa')
        if key and id not in key['sources']:raise HTTPException(403,'Origem não permitida para esta chave')
    def idempotent(c,request,u,body):
        key=request.headers.get('idempotency-key','')
        if not key or len(key)>200:raise HTTPException(422,'Idempotency-Key obrigatório, até 200 caracteres')
        digest=fingerprint({'path':request.url.path,'body':body})
        row=c.execute('SELECT hash,response FROM idempotency WHERE actor=? AND key=?',(u['id'],key)).fetchone()
        if row:
            if row[0]!=digest:raise HTTPException(409,'Chave idempotente reutilizada com conteúdo diferente')
            return key,digest,json.loads(row[1])
        return key,digest,None
    def remember(c,u,key,digest,result):c.execute('INSERT INTO idempotency VALUES (?,?,?,?)',(u['id'],key,digest,json.dumps(result,ensure_ascii=False)))
    def audit(c,u,action,id):store.event(c,{'id':uid(),'actor_id':u['id'],'action':action,'target_id':id,'at':now()})
    def heavy_job_capacity(c,user_id):
        active=[j for j in store.all(c,'job') if j['status'] in {'pending','preparing'}]
        active.extend(j for j in store.all(c,'import_job') if j['status'] in {'pending','processing'})
        if len(active)>=10 or sum(j['user_id']==user_id for j in active)>=2:
            raise HTTPException(429,'Fila ocupada: até dois trabalhos por usuário e dez no ambiente local',headers={'Retry-After':'30'})
    def owned_import(c,id,user,key):
        job=store.get(c,'import_job',id)
        if not job or job['user_id']!=user['id'] or (key and job.get('api_key_id')!=key['id']):
            raise HTTPException(404,'Importação não encontrada')
        return job
    def get_entity(c,id,collection=None):
        e=store.get(c,'entity',id)
        if not e or (collection is not None and (collection not in {'people','companies'} or e['entity_type']!=('person' if collection=='people' else 'company'))):raise HTTPException(404,'Cadastro não encontrado')
        return project(e)
    from .canonical_http import install_canonical_reads, maintain_canonical_cursors
    install_canonical_reads(app, canonical_reads, store=store, security=security, auth=auth, audit=audit)
    from .canonical_enrichment import install_canonical_writes
    install_canonical_writes(app, canonical_reads, store=store, auth=auth, source_check=source_check)
    from .canonical_validation import install_canonical_validation
    install_canonical_validation(app, canonical_reads, store=store, auth=auth, source_check=source_check)
    from .canonical_scalar_patch import install_canonical_scalar_patch
    install_canonical_scalar_patch(app, canonical_reads, store=store, auth=auth, source_check=source_check)
    from .canonical_catalog import install_canonical_catalog
    install_canonical_catalog(app, canonical_reads, store=store, security=security, auth=auth)
    @app.get('/api/v1/health')
    def health():return {'status':'ok','environment':'isolated-development','production_connected':False}
    @app.post('/api/v1/auth/activate')
    def activate(body:dict):
        if not isinstance(body.get('token'),str) or not isinstance(body.get('password'),str):raise HTTPException(422,'Informe convite e nova senha')
        with store.transaction() as c:
            result=security.activate(c,body['token'],body['password']);ch=store.get(c,'challenge',hashed(result['challenge']));audit(c,{'id':ch['user_id']},'activate_account',ch['user_id']);return result
    @app.post('/api/v1/auth/login')
    def login(body:Login):
        account_limit('password:'+body.username.strip().casefold())
        with store.transaction() as c:return security.login(c,body.username,body.password)
    @app.post('/api/v1/auth/otp')
    def otp(body:OTP,response:Response):
        with store.transaction() as c:
            challenge_limit(c,body.challenge);token,result=security.complete(c,body.challenge,body.code);audit(c,result['user'],'login',result['user']['id'])
        response.set_cookie('bigbase_session',token,httponly=True,secure=not (testing or os.environ.get('BIGBASE_LOCAL_HTTP')=='1'),samesite='strict',max_age=43200)
        return result
    @app.get('/api/v1/auth/me')
    def me(request:Request):
        with store.transaction() as c:
            u,key=auth(c,request,'read');s=store.get(c,'session',hashed(request.cookies.get('bigbase_session','')))
            return {'user':security.public(u),'csrf':s['csrf'] if s else None}
    @app.post('/api/v1/auth/step-up')
    def step_up(body:StepUpOTP,request:Request):
        with store.transaction() as c:
            u,key=auth(c,request,'admin')
            security.human_session(c,request,u,key)
            account_limit('otp:'+u['id'])
            result=security.step_up(c,request,u,key,body.code)
            audit(c,u,'admin_step_up',u['id'])
            return result
    @app.post('/api/v1/auth/recovery')
    def recover(body:dict):
        if not isinstance(body.get('challenge'),str) or not isinstance(body.get('code'),str):raise HTTPException(422,'Informe desafio e código de recuperação')
        with store.transaction() as c:
            challenge_limit(c,body['challenge']);return security.recover(c,body['challenge'],body['code'])
    @app.post('/api/v1/auth/logout')
    def logout(request:Request,response:Response):
        with store.transaction() as c:
            auth(c,request,'read');c.execute('DELETE FROM objects WHERE kind=? AND id=?',('session',hashed(request.cookies.get('bigbase_session',''))))
        response.delete_cookie('bigbase_session');return {'ok':True}
    @app.get('/api/v1/saved-searches')
    def saved_searches(request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'read');return {'items':[x for x in store.all(c,'saved_search') if x['user_id']==u['id'] and x['active']]}
    @app.post('/api/v1/saved-searches',status_code=201)
    def save_search(body:dict,request:Request):
        validate_filter(body.get('filters',{}))
        if not isinstance(body.get('name'),str) or not body['name'].strip():raise HTTPException(422,'Dê um nome à pesquisa')
        if body.get('entity_type','person') not in {'person','company'}:raise HTTPException(422,'Tipo inválido')
        applied_sort=resolve_sort(body)
        if type(body.get('include_invalid',False)) is not bool:raise HTTPException(422,'include_invalid deve ser booleano')
        with store.transaction() as c:
            u,_=auth(c,request,'read');ik,digest,cached=idempotent(c,request,u,body)
            if cached:return cached
            record={'id':uid(),'user_id':u['id'],'name':body['name'].strip()[:160],'filters':body.get('filters',{}),'entity_type':body.get('entity_type','person'),**({'sorts':body['sorts']} if 'sorts' in body else {'sort':body.get('sort','name'),'direction':body.get('direction','asc')}),'applied_sort':applied_sort,'include_invalid':body.get('include_invalid',False),'active':True,'created_at':now()}
            store.put(c,'saved_search',record);remember(c,u,ik,digest,record);audit(c,u,'save_search',record['id']);return record
    @app.patch('/api/v1/saved-searches/{id}')
    def archive_search(id:str,request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'read');record=store.get(c,'saved_search',id)
            if not record or record['user_id']!=u['id']:raise HTTPException(404)
            record.update(active=False,archived_at=now());store.put(c,'saved_search',record);audit(c,u,'archive_search',id);return {'ok':True}
    @app.post('/api/v1/imports',status_code=202)
    def create_import(body:dict,request:Request):
        if set(body)-{'name','entries'}:raise HTTPException(422,'Propriedades desconhecidas no pedido de importação')
        with store.transaction() as c:
            user,key=auth(c,request,'enrich');ik,digest,cached=idempotent(c,request,user,body)
            if cached:
                owned_import(c,cached['id'],user,key)
                return cached
            heavy_job_capacity(c,user['id'])
            job=imports.prepare(body.get('entries'),user['id'],body.get('name','Importação de cadastros'),key['id'] if key else None)
            store.put(c,'import_job',job);result=imports.public(job)
            remember(c,user,ik,digest,result);audit(c,user,'create_import',job['id'])
        executor.submit(imports.run,job['id'])
        return result
    @app.get('/api/v1/imports')
    def list_imports(request:Request):
        with store.transaction() as c:
            user,key=auth(c,request,'enrich')
            rows=[imports.public(j) for j in store.all(c,'import_job') if j['user_id']==user['id'] and (not key or j.get('api_key_id')==key['id'])]
            return {'items':sorted(rows,key=lambda j:(j['created_at'],j['id']),reverse=True)}
    @app.get('/api/v1/imports/{id}')
    def import_detail(id:str,request:Request,offset:int=0,limit:int=25):
        if not 0<=offset<=1000 or not 1<=limit<=100:raise HTTPException(422,'Paginação inválida no adaptador local')
        with store.transaction() as c:
            user,key=auth(c,request,'enrich');job=owned_import(c,id,user,key)
            detail=imports.result_detail(job);rows=detail['items']
            return {'job':imports.public(job),'results':rows[offset:offset+limit],'offset':offset,'limit':limit,'total_results':len(rows),'complete':offset+limit>=len(rows)}
    @app.post('/api/v1/imports/{id}/cancel')
    def cancel_import(id:str,request:Request):
        with store.transaction() as c:
            user,key=auth(c,request,'enrich');job=owned_import(c,id,user,key)
            if job['status'] not in {'pending','processing'}:raise HTTPException(409,'Importação já finalizada')
            job.update(status='cancelled',updated_at=now(),completed_at=now())
            job['cancelled_at']=job['completed_at']
            store.put(c,'import_job',job);audit(c,user,'cancel_import',id)
            return imports.public(job)
    @app.post('/api/v1/imports/{id}/resume',status_code=202)
    def resume_import(id:str,body:dict,request:Request):
        if body:raise HTTPException(422,'A retomada não altera as entradas do trabalho')
        with store.transaction() as c:
            user,key=auth(c,request,'enrich');job=owned_import(c,id,user,key)
            ik,digest,cached=idempotent(c,request,user,body)
            if cached:return cached
            if job['status'] not in {'failed','cancelled'} or job['processed']>=job['total']:
                raise HTTPException(409,'Não há entradas pendentes neste trabalho interrompido')
            if (job.get('error') or {}).get('code')=='checkpoint_inconsistent':
                raise HTTPException(409,'Checkpoint exige revisão antes da retomada')
            heavy_job_capacity(c,user['id'])
            job.setdefault('resume_history',[]).append({'at':now(),'actor_id':user['id'],'previous_status':job['status'],'previous_error':job.get('error'),'processed':job['processed'],'previous_updated_at':job['updated_at'],'previous_failed_at':job.get('failed_at'),'previous_cancelled_at':job.get('cancelled_at'),'previous_completed_at':job.get('completed_at')})
            job.update(status='pending',updated_at=now(),error=None)
            for field in ('completed_at','failed_at','cancelled_at'):job.pop(field,None)
            store.put(c,'import_job',job);result=imports.public(job)
            remember(c,user,ik,digest,result);audit(c,user,'resume_import',id)
        executor.submit(imports.run,id)
        return result
    @app.get('/api/v1/search/catalog')
    def catalog(request:Request):
        with store.transaction() as c:auth(c,request,'read')
        return {'fields':sorted(FIELDS),'operators':sorted(OPS),'version':2,'sorting':sorting_catalog()}
    @app.get('/api/v1/stats')
    def stats(request:Request):
        with store.transaction() as c:
            auth(c,request,'read');es=store.all(c,'entity');jobs=store.all(c,'job')
            return {'people':sum(e['entity_type']=='person' for e in es),'companies':sum(e['entity_type']=='company' for e in es),'items':sum(len(e['items']) for e in es),'observations':sum(len(e['observations']) for e in es),'pending':sum(o.get('pending_reason') is not None for e in es for o in e['observations'])}
    @app.post('/api/v1/{collection}/enrich')
    def enrich(collection:str,body:Enrichment,request:Request):
        if collection not in {'people','companies'}:raise HTTPException(404)
        typ='person' if collection=='people' else 'company'
        if body.entity_type!=typ:raise HTTPException(422,'Tipo não corresponde à rota')
        with store.transaction() as c:
            u,key=auth(c,request,'enrich');source_check(c,body.source_id,key);ik,digest,cached=idempotent(c,request,u,body.model_dump())
            if cached:return cached
            entity=enrich_entity(store,c,body,typ,u['id']);audit(c,u,'enrich',entity['id']);remember(c,u,ik,digest,entity);return entity
    @app.post('/api/v1/{collection}/search')
    def search(collection:str,body:dict,request:Request):
        if collection not in {'people','companies'}:raise HTTPException(404)
        filters=body.get('filters',{});validate_filter(filters);limit=body.get('limit',50);offset=body.get('offset',0)
        if type(limit) is not int or not 1<=limit<=100 or type(offset) is not int or not 0<=offset<=10000:raise HTTPException(422,'Paginação inválida no adaptador local')
        applied_sort=resolve_sort(body)
        with store.transaction() as c:
            u,_=auth(c,request,'read');entities=[project(e) for e in store.all(c,'entity') if e['entity_type']==('person' if collection=='people' else 'company') and matches(e,filters,include_invalid=body.get('include_invalid',False))]
            entities=sort_entities(entities,applied_sort,include_invalid=body.get('include_invalid',False));audit(c,u,'search',collection)
            return {'items':entities[offset:offset+limit],'total':len(entities),'offset':offset,'limit':limit,'complete':offset+limit>=len(entities),'search_backend':'development-adapter','applied_sort':applied_sort}
    @app.get('/api/v1/{collection}/{id}/relationships')
    def relationships(collection:str,id:str,request:Request):
        if collection not in {'people','companies'}:raise HTTPException(404)
        with store.transaction() as c:
            u,_=auth(c,request,'read');entity=get_entity(c,id)
            if entity['entity_type']!=('person' if collection=='people' else 'company'):raise HTTPException(404)
            audit(c,u,'relationships',id);return connections(entity,store.all(c,'entity'))
    @app.get('/api/v1/{collection}/{id}/history')
    def history(collection:str,id:str,request:Request):
        with store.transaction() as c:auth(c,request,'read');return {'items':get_entity(c,id,collection)['observations']}
    @app.get('/api/v1/people/{id}')
    @app.get('/api/v1/companies/{id}')
    def detail(id:str,request:Request):
        with store.transaction() as c:u,_=auth(c,request,'read');e=get_entity(c,id,request.url.path.split('/')[3]);audit(c,u,'read',id);return e
    @app.patch('/api/v1/{collection}/{id}/items/{item_id}')
    def patch_item(collection:str,id:str,item_id:str,body:dict,request:Request):
        if set(body)-{'source_id','observed_at','source_updated_at','reason','flags','flag_evidence'}:raise HTTPException(422,'Para novo valor, enriqueça o cadastro e preserve o anterior')
        if collection not in {'people','companies'}:raise HTTPException(404)
        try:inp=ItemInput(kind='custom',value={},flags=body.get('flags',{}),flag_evidence=body.get('flag_evidence',{}));timestamp(body.get('observed_at'));timestamp(body.get('source_updated_at'))
        except (ValueError,TypeError):raise HTTPException(422,'Flags ou data inválidos')
        with store.transaction() as c:
            u,key=auth(c,request,'validate');source_check(c,body.get('source_id'),key);entity=get_entity(c,id,collection)
            item=next((x for x in entity['items'] if x['id']==item_id),None)
            if not item:raise HTTPException(404,'Item não encontrado')
            version=request.headers.get('if-match')
            if version is None:raise HTTPException(428,'If-Match obrigatório')
            if version.strip('"')!=str(item['version']):raise HTTPException(409,'Item alterado por outra operação')
            if not inp.flags and not inp.flag_evidence:raise HTTPException(422,'Informe ao menos uma flag')
            before=len(entity['observations'])
            record_flags(entity,item,inp,body['source_id'],body.get('source_updated_at') or body.get('observed_at'),u['id'],uid(),body.get('reason',''),now())
            for obs in entity['observations'][before:]:obs.update(source_updated_at=body.get('source_updated_at'),source_observed_at=body.get('observed_at'))
            if item['kind']=='custom':
                metadata={'field_id':item.get('field_id',item['value'].get('field_id')),'field_definition_version':item.get('field_definition_version'),'classification_state':item.get('classification_state','pending'),'value_policy':VALUE_POLICY}
                bind_definition(item,entity['observations'][before:],metadata)
            item['version']+=1;entity['version']+=1;entity['updated_at']=now();project(entity);store.put(c,'entity',entity);audit(c,u,'validate',item_id);return entity
    @app.get('/api/v1/admin/{kind}')
    def list_admin(kind:str,request:Request):
        types={'users':'user','sources':'source','fields':'field','api-keys':'api_key'}
        if kind not in types:raise HTTPException(404)
        with store.transaction() as c:
            user,key=auth(c,request,'read' if kind in {'sources','fields','api-keys'} else 'admin');rows=store.all(c,types[kind])
            if kind=='users':rows=[security.public(x) for x in rows]
            if kind=='fields':rows=[public_definition(x) for x in rows]
            if kind=='api-keys':
                security.human_session(c,request,user,key)
                rows=[security.public_api_key(x) for x in rows if x['user_id']==user['id'] or 'admin' in user['permissions']]
            return {'items':rows}
    @app.post('/api/v1/admin/invitations',status_code=201)
    def invitation(body:dict,request:Request):
        if not isinstance(body.get('username'),str):raise HTTPException(422,'Informe usuário')
        permissions=body.get('permissions',['read'])
        if not isinstance(permissions,list) or any(x not in PERMISSIONS or x=='admin' for x in permissions):raise HTTPException(422,'Permissões do usuário normal inválidas')
        with store.transaction() as c:
            u,key=auth(c,request,'admin');security.require_recent_totp(c,request,u,key)
            result=security.invite(c,body['username'],permissions);audit(c,u,'invite_user',result['user']['id']);return result
    @app.post('/api/v1/admin/users')
    def create_user(body:dict,request:Request):
        with store.transaction() as c:
            u,key=auth(c,request,'admin');security.require_recent_totp(c,request,u,key)
            new=security.create_user(c,body.get('username',''),body.get('password',''),body.get('role','user'),body.get('permissions'));audit(c,u,'create_user',new['id']);return security.public(new)
    @app.patch('/api/v1/admin/users/{id}')
    def disable_user(id:str,body:dict,request:Request):
        with store.transaction() as c:
            u,key=auth(c,request,'admin');security.require_recent_totp(c,request,u,key)
            target=store.get(c,'user',id)
            if not target:raise HTTPException(404)
            if id==u['id']:raise HTTPException(409,'Não desative seu próprio acesso')
            if type(body.get('active')) is not bool:raise HTTPException(422,'Informe active true/false')
            if target.get('activation_pending') and body['active']:raise HTTPException(409,'O usuário precisa concluir o convite')
            if not body['active']:security.revoke_access(c,target['id'])
            target['active']=body['active'];store.put(c,'user',target);audit(c,u,'user_status',id);return security.public(target)
    @app.post('/api/v1/admin/api-keys')
    def create_key(body:dict,request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'admin')
            if not security.check_code(u,body.get('otp','')):raise HTTPException(401,'Confirme com um novo código OTP')
            store.put(c,'user',u);target=store.get(c,'user',body.get('user_id',u['id']));scopes=body.get('scopes',['read']);sources=body.get('sources',['manual'])
            if not isinstance(scopes,list) or not scopes or any(not isinstance(x,str) for x in scopes) or not isinstance(sources,list) or any(not isinstance(x,str) for x in sources):raise HTTPException(422,'Escopos e origens devem ser listas de identificadores')
            if not target or not target['active'] or not target['otp_enabled'] or set(scopes)-set(target['permissions']):raise HTTPException(422,'Usuário ou escopos inválidos')
            for source in sources:source_check(c,source,None)
            token='bb_'+secrets.token_urlsafe(32);record={'id':hashed(token),'public_id':uid(),'user_id':target['id'],'name':str(body.get('name','Integração'))[:120],'scopes':scopes,'sources':sources,'active':True,'created_at':now(),'expires_at':(datetime.now(timezone.utc)+timedelta(days=90)).isoformat()}
            store.put(c,'api_key',record);audit(c,u,'create_key',record['public_id']);return {'id':record['public_id'],'key':token,'expires_at':record['expires_at']}
    @app.patch('/api/v1/admin/api-keys/{id}')
    def revoke_key(id:str,request:Request):
        with store.transaction() as c:
            u,auth_key=auth(c,request,'read');security.human_session(c,request,u,auth_key)
            key=security.owned_api_key(c,u,id)
            result=security.revoke_key_chain(c,u,key);audit(c,u,'revoke_key_chain',id);return result
    @app.post('/api/v1/admin/api-keys/{id}/rotate',status_code=201)
    def rotate_key(id:str,body:dict,request:Request):
        with store.transaction() as c:
            u,key=auth(c,request,'read');security.human_session(c,request,u,key)
            account_limit('otp:'+u['id'])
            result,_=security.rotate_api_key(c,request,u,key,id,body,request.headers.get('idempotency-key'))
            return result
    @app.post('/api/v1/admin/fields')
    def create_field(body:FieldDefinitionInput,request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'admin');record=create_definition(store,c,body,u['id']);audit(c,u,'create_field',record['id']);return record
    @app.patch('/api/v1/admin/fields/{id}')
    def update_field(id:str,body:dict,request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'admin');record=update_definition(store,c,id,body,request.headers.get('if-match'),u['id']);audit(c,u,'update_field',id);return record
    @app.post('/api/v1/admin/{kind}')
    def create_catalog(kind:str,body:dict,request:Request):
        if kind!='sources':raise HTTPException(404)
        with store.transaction() as c:
            u,_=auth(c,request,'admin');id=str(body.get('id',uid()));typ='source'
            if store.get(c,typ,id):raise HTTPException(409,'Identificador já cadastrado')
            if not body.get('name'):raise HTTPException(422,'Nome obrigatório')
            record={'id':id,'name':str(body['name'])[:160],'active':True,'created_at':now(),'version':1}
            store.put(c,typ,record);audit(c,u,'create_'+typ,id);return record
    @app.get('/api/v1/audit')
    def audit_list(request:Request):
        with store.transaction() as c:auth(c,request,'admin');return {'items':[json.loads(x[0]) for x in c.execute('SELECT body FROM events ORDER BY seq DESC LIMIT 200')]}
    @app.post('/api/v1/bulk-queries',status_code=202)
    def bulk(body:dict,request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'export');ik,digest,cached=idempotent(c,request,u,body)
            if cached:return cached
            heavy_job_capacity(c,u['id'])
            if body.get('upload_id'):
                upload=store.get(c,'upload',body['upload_id'])
                if not upload or upload['user_id']!=u['id']:raise HTTPException(404,'Lista não encontrada')
                if body.get('text') or body.get('entries'):raise HTTPException(422,'Use uma lista enviada ou texto/entradas por trabalho')
                body={**body,'entries':[{'input_id':str(n+1),'field':body.get('input_field','name'),'value':v,'mode':body.get('match_mode','eq'),'document_type':body.get('document_type','CPF')} for n,v in enumerate(upload['values'])]}
            job=prepare_job(body,store.all(c,'entity'),u['id']);store.put(c,'job',job);result=exports.public(job);remember(c,u,ik,digest,result);audit(c,u,'bulk_query',job['id'])
        executor.submit(exports.run,job['id']);return result
    @app.post('/api/v1/bulk-queries/uploads',status_code=201)
    def upload_list(file:UploadFile,request:Request):
        with store.transaction() as c:u,_=auth(c,request,'export')
        data=file.file.read(2*1024*1024+1)
        if len(data)>2*1024*1024:raise HTTPException(413,'Arquivo excede limite local de 2 MiB')
        name=(file.filename or '').lower();values=[]
        try:
            if name.endswith('.csv'):
                values=[cell.strip() for row in csv.reader(io.StringIO(data.decode('utf-8-sig'))) for cell in row if cell.strip()]
            elif name.endswith('.xlsx'):
                from openpyxl import load_workbook
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    if sum(x.file_size for x in archive.infolist())>32*1024*1024:raise ValueError('Arquivo expandido excede 32 MiB')
                book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
                try:
                    for sheet in book:
                        if (sheet.max_row or 0)*(sheet.max_column or 0)>200000:raise ValueError('Dimensões acima do limite do importador local')
                        for row in sheet.iter_rows(values_only=True):
                            for value in row:
                                if value is not None and str(value).strip():values.append(str(value).strip())
                            if len(values)>100000:raise ValueError('Mais de 100.000 entradas')
                finally:book.close()
            else:raise ValueError('Use CSV UTF-8 ou XLSX')
        except Exception as exc:raise HTTPException(422,'Lista inválida: '+str(exc))
        if not values or len(values)>100000:raise HTTPException(422,'Lista vazia ou acima de 100.000 entradas')
        with store.transaction() as c:
            u,_=auth(c,request,'export');upload={'id':uid(),'user_id':u['id'],'values':values,'created_at':now()};store.put(c,'upload',upload);audit(c,u,'upload_query_list',upload['id'])
        return {'id':upload['id'],'count':len(values),'preview':values[:10]}
    def owned_job(c,id,u):
        j=store.get(c,'job',id)
        if not j or j['user_id']!=u['id']:raise HTTPException(404,'Trabalho não encontrado')
        return j
    @app.get('/api/v1/bulk-queries')
    def jobs(request:Request):
        with store.transaction() as c:u,_=auth(c,request,'export');return {'items':sorted([exports.public(j) for j in store.all(c,'job') if j['user_id']==u['id']],key=lambda j:j['created_at'])}
    @app.get('/api/v1/bulk-queries/{id}')
    def job(id:str,request:Request):
        with store.transaction() as c:u,_=auth(c,request,'export');return exports.public(owned_job(c,id,u))
    @app.post('/api/v1/bulk-queries/{id}/cancel')
    def cancel(id:str,request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'export');j=owned_job(c,id,u)
            if j['status'] not in {'pending','preparing'}:raise HTTPException(409,'Trabalho já finalizado')
            j.update(status='cancelled',updated_at=now());store.put(c,'job',j);audit(c,u,'cancel_job',id);return exports.public(j)
    @app.get('/api/v1/bulk-queries/{id}/files/result')
    def download(id:str,request:Request):
        with store.transaction() as c:
            u,_=auth(c,request,'export');j=owned_job(c,id,u)
            if j['status']!='completed':raise HTTPException(409,'Arquivo ainda não disponível')
            if timestamp(j['expires_at'])<datetime.now(timezone.utc):raise HTTPException(410,'Arquivo expirado')
            audit(c,u,'download',id)
        artifact=j.get('artifact','xlsx');media='application/zip' if artifact=='zip' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return FileResponse(root/'exports'/(id+'.'+artifact),filename='BIG_BASE_'+id[:8]+'.'+artifact,media_type=media)
    frontend=Path(__file__).resolve().parents[2]/'frontend'/'dist'
    if frontend.exists():app.mount('/',StaticFiles(directory=frontend,html=True),name='panel')
    return app
