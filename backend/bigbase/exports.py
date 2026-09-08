from copy import deepcopy
import csv
import io
import json
import hashlib
import math
import sys
import zipfile
import tempfile
from pathlib import Path
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from .domain import now, uid, canonical, normalize, FLAGS, project, timestamp
from .search import matches, validate_filter
from .sorting import resolve_sort, sort_entities
from fastapi import HTTPException

SHEETS={'identity':'Identidade','document':'Documentos','phone':'Telefones','email':'Emails','address':'Enderecos','username':'Usernames','relationship':'Relacoes','activity':'Atividades','custom':'Campos adicionais'}

def exact_number_text(value):
    """Return exact JSON text when Excel's 15-digit numeric representation is unsafe."""
    if type(value) not in {int,float}:return None
    if isinstance(value,int):return canonical(value) if abs(value)>999999999999999 else None
    if not math.isfinite(value):raise ValueError('Número não finito não pode ser exportado')
    if (abs(value)>999999999999999
        or (value!=0 and abs(value)<sys.float_info.min)
        or (value==0 and math.copysign(1,value)<0)
        or float(format(value,'.15g'))!=value):
        return canonical(value)
    return None

def parse_entries(body):
    entries=body.get('entries',[])
    if body.get('text'):
        entries=[{'input_id':str(n+1),'field':body.get('input_field','name'),'value':cell.strip(),'mode':body.get('match_mode','eq'),'document_type':body.get('document_type','CPF')} for n,cell in enumerate(c for row in csv.reader(io.StringIO(body['text'])) for c in row) if cell.strip()]
    if not isinstance(entries,list) or len(entries)>100000: raise HTTPException(422,'Lista inválida; limite local de 100.000 entradas')
    parsed=[]
    for n,e in enumerate(entries):
        if not isinstance(e,dict): raise HTTPException(422,'Entrada deve ser objeto')
        field=e.get('field','name'); value=e.get('value'); error=None
        if field=='document':
            document=normalize('document',{'type':e.get('document_type','CPF'),'number':value})[0]
            value=document['number']
            if document.get('syntax_valid') is False:error='Documento com formato ou dígito verificador inválido'
            node={'item':{'and':[{'field':'kind','value':'document'},{'field':'type','value':document['type']},{'field':'number','value':value},{'field':'country','value':document['country']}]}}
        elif field=='phone':
            phone=normalize('phone',{'number':value,'country':e.get('country','BR')})[0];value=phone['number']
            node={'item':{'and':[{'field':'kind','value':'phone'},{'field':'number','value':value}]}}
        else: node={'field':field,'op':e.get('mode','eq'),'value':value}
        try: validate_filter(node)
        except HTTPException as exc: error=exc.detail
        if value is None or value=='': error='Entrada vazia'
        if e.get('filters'): validate_filter(e['filters']); node={'and':[node,e['filters']]}
        parsed.append({'input_id':str(e.get('input_id',n+1)),'value':e.get('value'),'filter':node,'error':error})
    if len({e['input_id'] for e in parsed})!=len(parsed):raise HTTPException(422,'Cada input_id deve ser único; valores repetidos podem ter IDs diferentes')
    return parsed

def prepare_job(body,entities,owner):
    if body.get('format','xlsx')!='xlsx':raise HTTPException(422,'Esta modalidade entrega XLSX')
    body=deepcopy(body)
    applied_sort=resolve_sort(body,default_legacy=False)
    filters=body.get('filters',{}); validate_filter(filters); entries=parse_entries(body)
    if not entries and not filters and not body.get('all_records'): raise HTTPException(422,'Informe lista, filtros ou seleção explícita de todos')
    typ=body.get('entity_type','person')
    if typ not in {'person','company'}: raise HTTPException(422,'Tipo de entidade inválido')
    return {'id':uid(),'user_id':owner,'name':str(body.get('name','Consulta em massa'))[:160],'status':'pending','progress_percent':0,'phase':'Na fila','created_at':now(),'updated_at':now(),'cutoff':now(),'criteria':body,'applied_sort':applied_sort,'entries':entries,'snapshot':[project(e) for e in entities if e['entity_type']==typ], 'entity_count':0,'error':None}

class Exporter:
    def __init__(self,store,root,row_limit=1048576,volume_limit=50000,column_limit=16384):
        self.store=store; self.root=Path(root); self.row_limit=row_limit;self.volume_limit=volume_limit;self.column_limit=column_limit
        if row_limit<2 or volume_limit<1 or column_limit<16:raise ValueError('Limites de exportação inválidos')
    def cleanup_expired(self):
        removed=0
        with self.store.transaction() as c:
            for job in self.store.all(c,'job'):
                if job['status']!='completed' or job.get('artifact_expired') or not job.get('expires_at') or timestamp(job['expires_at'])>datetime.now(timezone.utc):continue
                artifact=job.get('artifact','xlsx')
                if artifact not in {'xlsx','zip'}:continue
                (self.root/(job['id']+'.'+artifact)).unlink(missing_ok=True)
                job.update(artifact_expired=True,artifact_removed_at=now());self.store.put(c,'job',job);removed+=1
        return removed

    def public(self,j): return {k:v for k,v in j.items() if k not in {'snapshot','entries'}}
    def definition_versions(self,j,selected):
        field_ids={value for e in selected for item in e['items'] for value in [item.get('field_id'),item.get('value',{}).get('field_id')] if isinstance(value,str)}
        field_ids.update(o['field_id'] for e in selected for o in e['observations'] if isinstance(o.get('field_id'),str))
        field_ids.update(s['field_id'] for e in selected for item in e['items'] for s in item['fields'].values() if isinstance(s.get('field_id'),str))
        versions={}
        with self.store.transaction() as c:
            # Immutable versions retain the definition that originally interpreted an observation.
            for record in self.store.all(c,'field_definition_version'):
                if record.get('field_id') in field_ids:
                    definition=record['definition'];versions[(record['field_id'],definition['version'])]=definition
            # Legacy definitions may predate the separate immutable-version registry.
            for record in self.store.all(c,'field'):
                if record['id'] not in field_ids:continue
                for definition in [*record.get('definition_history',[]),{k:v for k,v in record.items() if k!='definition_history'}]:
                    versions.setdefault((record['id'],definition.get('version',1)),definition)
        cutoff=timestamp(j['cutoff']);result=[]
        for (field_id,version),definition in sorted(versions.items()):
            changed_at=definition.get('updated_at') or definition.get('created_at')
            if changed_at and timestamp(changed_at)>cutoff:continue
            result.append({'field_id':field_id,'field_definition_version':version,**{k:v for k,v in definition.items() if k not in {'id','version','definition_history'}}})
        return result
    def update(self,j,**values):
        with self.store.transaction() as c:
            current=self.store.get(c,'job',j['id'])
            if current and current['status']=='cancelled': raise InterruptedError('Cancelado')
            j.update(values,updated_at=now()); self.store.put(c,'job',j)
    def run(self,id):
        with self.store.transaction() as c: j=self.store.get(c,'job',id)
        if not j or j['status'] not in {'pending','preparing'}: return
        path=self.root/(id+'.xlsx'); partial=path.with_suffix('.partial.xlsx')
        try:
            self.update(j,status='preparing',phase='Selecionando correspondências',progress_percent=5)
            filters=j['criteria'].get('filters',{}); selected=[]; links=[]; counts={e['input_id']:0 for e in j['entries']}
            for index,entity in enumerate(j['snapshot']):
                if not matches(entity,filters,include_invalid=j['criteria'].get('include_invalid',False)): continue
                hits=[e for e in j['entries'] if not e['error'] and matches(entity,e['filter'],include_invalid=j['criteria'].get('include_invalid',False))]
                if hits or not j['entries']:
                    selected.append(entity)
                    for e in hits: links.append([e['input_id'],entity['id']]); counts[e['input_id']]+=1
                if index%100==0:self.update(j,progress_percent=5+int(20*(index+1)/max(1,len(j['snapshot']))))
            applied_sort=j.get('applied_sort') or resolve_sort(j['criteria'],default_legacy=False)
            selected=sort_entities(selected,applied_sort,include_invalid=j['criteria'].get('include_invalid',False))
            self.update(j,applied_sort=applied_sort,phase='Gerando planilhas',progress_percent=30,entity_count=len(selected))
            definitions=self.definition_versions(j,selected)
            self.root.mkdir(parents=True,exist_ok=True)
            volumes=[];sheet_counts={}
            with tempfile.TemporaryDirectory(prefix=id+'-',dir=self.root) as temporary:
                chunks=[selected[n:n+self.volume_limit] for n in range(0,len(selected),self.volume_limit)] or [[]]
                for number,chunk in enumerate(chunks,1):
                    chunk_ids={e['id'] for e in chunk}
                    part=Path(temporary)/('volume_%04d.xlsx'%number)
                    stats=self.write_volume(j,chunk,[link for link in links if link[1] in chunk_ids],counts,part,30+int(60*(number-1)/len(chunks)),60/len(chunks),definitions)
                    with part.open('rb') as stream:digest=hashlib.file_digest(stream,'sha256').hexdigest()
                    volumes.append({'file':part.name,'sha256':digest,'entities':len(chunk),'sheet_counts':stats})
                    sheet_counts.update({(str(number)+':'+k if len(chunks)>1 else k):v for k,v in stats.items()})
                if len(volumes)==1:
                    artifact='xlsx';Path(temporary,volumes[0]['file']).replace(partial)
                else:
                    artifact='zip';path=self.root/(id+'.zip');partial=path.with_suffix('.partial.zip')
                    manifest={'job_id':id,'cutoff':j['cutoff'],'entity_count':len(selected),'criteria':j['criteria'],'applied_sort':j['applied_sort'],'volumes':volumes}
                    with zipfile.ZipFile(partial,'w',compression=zipfile.ZIP_STORED,allowZip64=True) as archive:
                        for volume in volumes:archive.write(Path(temporary)/volume['file'],volume['file'])
                        archive.writestr('manifesto.json',json.dumps(manifest,ensure_ascii=False,indent=2))
                    with zipfile.ZipFile(partial) as archive:
                        if archive.testzip() is not None:raise RuntimeError('Falha de integridade do ZIP')
                self.update(j,phase='Publicando arquivo verificado',progress_percent=99)
                partial.replace(path)
                with path.open('rb') as stream:digest=hashlib.file_digest(stream,'sha256').hexdigest()
                self.update(j,status='completed',progress_percent=100,phase='Disponível para download',artifact=artifact,sha256=digest,sheet_counts=sheet_counts,volumes=volumes,completed_at=now(),expires_at=(datetime.now(timezone.utc)+timedelta(days=7)).isoformat(),snapshot=[])
        except InterruptedError:
            partial.unlink(missing_ok=True);path.unlink(missing_ok=True)
        except Exception as exc:
            partial.unlink(missing_ok=True);path.unlink(missing_ok=True)
            with self.store.transaction() as c:
                current=self.store.get(c,'job',id)
                if current['status']!='cancelled':current.update(status='failed',error=str(exc),updated_at=now());self.store.put(c,'job',current)

    def write_volume(self,j,selected,links,counts,partial,base_progress,span_progress,definitions=()):
        wb=Workbook(write_only=True); sheet_counts={}; sheets={}; row_counts={}; parts={}; fragments=[]; structures=[]
        exact_numbers={}; validation_hashes={}
        def content_signature(value,exact=None,metadata=False):
            if exact:return ['exact_number',exact[0],value]
            if type(value) is bool:return ['boolean',value]
            if type(value) in {int,float}:return ['number',float(value).hex()]
            if metadata and isinstance(value,str):return ['text',value]
            return None
        def add_content_digest(digest,row,column,signature):
            if signature is not None:digest.update(canonical([row,column,signature]).encode('utf-8')+b'\n')
        def flatten_structure(ref,value,path=''):
            if isinstance(value,dict):
                structures.append([ref,path,'object',None])
                for k,v in value.items():flatten_structure(ref,v,path+'/'+str(k).replace('~','~0').replace('/','~1'))
            elif isinstance(value,list):
                structures.append([ref,path,'array',None])
                for n,v in enumerate(value):flatten_structure(ref,v,path+'/'+str(n))
            else:structures.append([ref,path,'null' if value is None else 'boolean' if isinstance(value,bool) else 'number' if isinstance(value,(float,int)) else 'text',value])
        def cell(ws,value,row,column):
            if value is None:return None
            if isinstance(value,(dict,list)):
                ref=uid();flatten_structure(ref,value);value='structure:'+ref
            number_text=exact_number_text(value)
            exact=None
            if number_text is not None:
                exact=('integer' if type(value) is int else 'float',number_text)
                exact_numbers[(ws.title,row,column)]=exact
                value=number_text
            c=WriteOnlyCell(ws,value=value)
            if isinstance(value,str):
                # Always explicit text: untrusted input must never become an Excel formula.
                c.data_type='s'
                value=''.join(ch if ord(ch)>=32 or ch in '\t\n\r' else '\\u%04x'%ord(ch) for ch in value)
                if len(value)>32767:
                    ref=uid(); fragments.extend([[ref,n//30000,value[n:n+30000]] for n in range(0,len(value),30000)]);value='fragment:'+ref
                c.value=value; c.data_type='s'
            add_content_digest(validation_hashes[ws.title],row,column,content_signature(value,exact,ws.title.startswith('Numeros exatos')))
            return c
        def write(name,headers,row):
            if name not in sheets or row_counts[name]>=self.row_limit:
                parts[name]=parts.get(name,0)+1; title=name if parts[name]==1 else name[:25]+'_'+str(parts[name])
                ws=wb.create_sheet(title);validation_hashes[title]=hashlib.sha256();ws.freeze_panes='A2';ws.append([cell(ws,h,1,col) for col,h in enumerate(headers,1)]);sheets[name]=ws;row_counts[name]=1;sheet_counts[title]=0
            ws=sheets[name];ws.append([cell(ws,x,row_counts[name]+1,col) for col,x in enumerate(row,1)]);row_counts[name]+=1;sheet_counts[ws.title]+=1
        headers=['entidade_id','tipo','nome','versao','criado_em','atualizado_em']
        kind_columns={kind:sorted({k for ent in selected for item in ent['items'] if item['kind']==kind for k in item['value']}) for kind in SHEETS}
        item_metadata_columns=sorted({key for entity in selected for item in entity['items'] for key in item if key not in {'id','value','flags','fields','flag_details'}})
        schema_columns=['field_id','field_definition_version','classification_state']
        for idx,e in enumerate(selected):
            write('Cadastros',headers,[e['id'],e['entity_type'],e['name'],e['version'],e['created_at'],e['updated_at']])
            for kind,title in SHEETS.items():
                columns=kind_columns[kind]
                for item in [x for x in e['items'] if x['kind']==kind]:
                    flags=sorted(FLAGS)
                    width=self.column_limit-2-len(flags)
                    for start in range(0,max(1,len(columns)),width):
                        part_columns=columns[start:start+width]
                        part_title=title if start==0 else title[:20]+'_col'+str(start//width+1)
                        write(part_title,['entidade_id','item_id',*part_columns,*flags],[e['id'],item['id'],*[item['value'].get(k) for k in part_columns],*[item['flags'].get(k) for k in flags]])
                    for field,state in item['fields'].items():
                        write('Estado dos campos',['entidade_id','item_id','campo','valor','origem','observado_em','recebido_em','observacao_id',*schema_columns],[e['id'],item['id'],field,state['value'],state['source_id'],state['observed_at'],state['received_at'],state['id'],*[state.get(key) for key in schema_columns]])
            for item in e['items']:
                for start in range(0,len(item_metadata_columns),self.column_limit-2):
                    columns=item_metadata_columns[start:start+self.column_limit-2]
                    title='Metadados dos itens' if start==0 else 'Metadados itens_col'+str(start//(self.column_limit-2)+1)
                    write(title,['entidade_id','item_id',*columns],[e['id'],item['id'],*[item.get(key) for key in columns]])
                for flag,details in item.get('flag_details',{}).items():
                    keys=['value','last_result','applicable','source_id','checked_at','received_at','expires_at','stale','method','reference','observation_id']
                    write('Confirmacoes',['entidade_id','item_id','flag',*keys],[e['id'],item['id'],flag,*[details.get(k) for k in keys]])
            for obs in e['observations']:
                keys=['id','item_id','path','value','input_value','previous','source_id','source_updated_at','source_observed_at','observed_at','received_at','external_id','actor_id','operation_id','reason','applied','pending_reason','verification','value_binding','normalization',*schema_columns]
                write('Historico',['entidade_id',*keys],[e['id'],*[obs.get(k) for k in keys]])
            if idx%25==0:self.update(j,progress_percent=base_progress+int(span_progress*.9*(idx+1)/max(1,len(selected))))
        for e in j['entries']:
            write('Entradas',['entrada_id','valor','estado','encontrados','erro'],[e['input_id'],e['value'],'invalid' if e['error'] else 'found' if counts[e['input_id']] else 'not_found',counts[e['input_id']],e['error']])
        for link in links:write('Correspondencias',['entrada_id','entidade_id'],link)
        definition_columns=['field_id','field_definition_version',*sorted({key for definition in definitions for key in definition}-{'field_id','field_definition_version'})]
        for definition in definitions:
            for start in range(0,len(definition_columns)-2,self.column_limit-2):
                columns=definition_columns[2+start:2+start+self.column_limit-2]
                title='Definicoes dos campos' if start==0 else 'Definicoes campos_col'+str(start//(self.column_limit-2)+1)
                write(title,[*definition_columns[:2],*columns],[definition['field_id'],definition['field_definition_version'],*[definition.get(key) for key in columns]])
        for priority,criterion in enumerate(j['applied_sort']['criteria'],1):
            write('Ordenacao',['contrato','versao','prioridade','campo','direcao','modo','ausentes','comparacao_texto'],[j['applied_sort']['contract'],j['applied_sort']['version'],priority,criterion['field'],criterion['direction'],criterion.get('mode'),j['applied_sort'].get('missing'),j['applied_sort'].get('text_comparison')])
        if not j['applied_sort']['criteria']:
            write('Ordenacao',['contrato','versao','prioridade','campo','direcao','modo','ausentes','comparacao_texto'],['snapshot_order',1,None,None,None,None,None,None])
        write('Resumo',['trabalho_id','data_corte','total_entidades','criterios'],[j['id'],j['cutoff'],len(selected),j['criteria']])
        for structured in structures:write('Valores estruturados',['referencia','caminho_json_pointer','tipo','valor'],structured)
        for f in fragments:write('Fragmentos',['referencia','posicao','conteudo'],f)
        for (title,row,column),(original_type,number_text) in exact_numbers.items():
            write('Numeros exatos',['aba','linha','coluna','tipo_json','tipo_original','valor_json'],[title,row,column,'number',original_type,number_text])
        for title,count in list(sheet_counts.items()):write('Manifesto',['aba','linhas'],[title,count])
        self.root.mkdir(parents=True,exist_ok=True);wb.save(partial)
        self.update(j,phase='Verificando arquivo',progress_percent=base_progress+int(span_progress*.95))
        check=load_workbook(partial,read_only=True,data_only=False)
        try:
            actual_ids=[row[0] for sheet in check if sheet.title=='Cadastros' or sheet.title.startswith('Cadastros_') for row in sheet.iter_rows(min_row=2,values_only=True)]
            if actual_ids!=[entity['id'] for entity in selected]:raise RuntimeError('Ordem ou IDs divergentes nos cadastros exportados')
            for title,count in sheet_counts.items():
                digest=hashlib.sha256();rows=0
                for rows,row in enumerate(check[title].iter_rows(),1):
                    for column,loaded_cell in enumerate(row,1):
                        exact=exact_numbers.get((title,rows,column))
                        if exact and (loaded_cell.data_type!='s' or loaded_cell.value!=exact[1]):raise RuntimeError('Conteúdo numérico divergente: '+title)
                        add_content_digest(digest,rows,column,content_signature(loaded_cell.value,exact,title.startswith('Numeros exatos')))
                if rows-1!=count:raise RuntimeError('Contagem divergente: '+title)
                if digest.digest()!=validation_hashes[title].digest():raise RuntimeError('Conteúdo numérico/booleano divergente: '+title)
        finally:check.close()
        return sheet_counts
