"""Synthetic regressions for value fidelity across search and XLSX publication."""
import copy
import json
import math

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from bigbase.domain import add_observation, now, project
from bigbase.exports import Exporter, prepare_job
from bigbase.search import matches, validate_filter
from bigbase.store import Store


def custom_entity(value):
    return {'items':[{'kind':'custom','value':{'field_id':'synthetic','value':value},'flags':{},'sources':[]}]}


@pytest.mark.parametrize(('actual','query','expected'),[
    (False,0,False), (True,1.0,False), (0,False,False), (1,True,False),
    (False,False,True), (None,None,True), (None,0,False),
    (0,0.0,True), (12,12.0,True), ('0',0,False),
    ([False,{'nested':True}],[0,{'nested':1}],False),
    ({'nested':[False,None,0]},{'nested':[False,None,0.0]},True),
    ({'nested':None},{},False),
    (9007199254740993,9007199254740992.0,False),
])
def test_search_eq_and_in_respect_recursive_json_types(actual,query,expected):
    entity=custom_entity(actual)
    for op,value in [('eq',query),('in',[query])]:
        node={'field':'value','op':op,'value':value}
        validate_filter(node)
        assert matches(entity,node) is expected


@pytest.mark.parametrize(('actual','bounds','expected'),[
    (False,[0,1],False), (True,[0,1],False), (False,[None,None],False),
    ([0],[None,None],False), ({'value':0},[None,None],False),
    (0,[0.0,1.0],True), (1.5,[1,2],True), (None,[0,1],False),
    ('01234567',['01000000','02000000'],True), ('1',[0,2],False),
])
def test_range_only_orders_compatible_numeric_or_text_values(actual,bounds,expected):
    node={'field':'value','op':'range','value':bounds}
    validate_filter(node)
    assert matches(custom_entity(actual),node) is expected


@pytest.mark.parametrize('bounds',[[False,1],[0,True],[None,False],[[0],None],[{'x':1},None],['0',1]])
def test_range_rejects_boolean_structured_or_mixed_bounds(bounds):
    with pytest.raises(HTTPException) as exc:
        validate_filter({'field':'value','op':'range','value':bounds})
    assert exc.value.status_code==422


def make_numeric_export(tmp_path):
    store=Store(tmp_path/'synthetic.sqlite')
    values={
        'field_id':'synthetic-numbers',
        'large_integer':9007199254740993,
        'negative_large_integer':-9007199254740993,
        'sixteen_digits':1234567890123456,
        'safe_integer':999999999999999,
        'precise_float':0.12345678901234566,
        'underflow_float':5e-324,
        'negative_zero':-0.0,
        'plain_float':1.5,
        'false_value':False,
        'zero_value':0,
        'null_value':None,
        'value':{'nested':[False,0,None,9007199254740993,0.12345678901234566]},
    }
    entity={'id':'synthetic-entity','entity_type':'person','version':1,'created_at':now(),'updated_at':now(),'items':[],'observations':[]}
    item={'id':'synthetic-item','kind':'custom','version':1,'fields':{},'sources':[]}
    entity['items'].append(item)
    for key,value in values.items():
        add_observation(entity,item,'value.'+key,value,'synthetic','2026-01-01T00:00:00Z','tester','synthetic-op','',now())
    add_observation(entity,item,'flag.valid',False,'synthetic','2026-01-01T00:00:00Z','tester','synthetic-op','',now())
    # Both the original input and previous value must survive history export.
    add_observation(entity,item,'value.large_integer',9007199254740995,'synthetic','2026-02-01T00:00:00Z','tester','next-synthetic-op','',now(),9007199254740997)
    project(entity)
    with store.transaction() as tx:
        job=prepare_job({'all_records':True},[entity],'synthetic-owner')
        store.put(tx,'job',job)
    exporter=Exporter(store,tmp_path/'exports',row_limit=12)
    return exporter,job,copy.deepcopy(entity)


def export_book(exporter,job):
    exporter.run(job['id'])
    with exporter.store.transaction() as tx:result=exporter.store.get(tx,'job',job['id'])
    assert result['status']=='completed',result.get('error')
    return load_workbook(exporter.root/(job['id']+'.xlsx'))


def numeric_metadata(book):
    result={}
    for title in book.sheetnames:
        if title.startswith('Numeros exatos'):
            for row in book[title].iter_rows(min_row=2,values_only=True):
                sheet,line,column,json_type,original_type,exact=row
                assert json_type=='number'
                cell=book[sheet].cell(line,column)
                assert cell.data_type=='s'
                assert cell.value==exact
                decoded=json.loads(exact)
                assert type(decoded) is (int if original_type=='integer' else float)
                result[(sheet,line,column)]=decoded
    return result


def decoded_cell(cell,metadata):
    return metadata.get((cell.parent.title,cell.row,cell.column),cell.value)


def test_xlsx_preserves_exact_numbers_in_values_history_and_nested_structures(tmp_path):
    exporter,job,entity=make_numeric_export(tmp_path)
    book=export_book(exporter,job)
    try:
        metadata=numeric_metadata(book)
        assert metadata
        assert any(title.startswith('Numeros exatos_') for title in book.sheetnames)
        item=entity['items'][0]
        sheet=book['Campos adicionais']
        columns={cell.value:cell.column for cell in sheet[1]}
        for field in ['large_integer','negative_large_integer','sixteen_digits','precise_float','underflow_float','negative_zero']:
            cell=sheet.cell(2,columns[field])
            actual=decoded_cell(cell,metadata)
            expected=item['value'][field]
            assert actual==expected and type(actual) is type(expected)
            if field=='negative_zero':assert math.copysign(1,actual)==-1
        assert sheet.cell(2,columns['false_value']).value is False
        assert sheet.cell(2,columns['zero_value']).value==0
        assert sheet.cell(2,columns['zero_value']).data_type=='n'
        assert sheet.cell(2,columns['plain_float']).value==1.5
        assert sheet.cell(2,columns['safe_integer']).value==999999999999999
        reference=sheet.cell(2,columns['value']).value.removeprefix('structure:')
        nested={}
        for title in book.sheetnames:
            if title.startswith('Valores estruturados'):
                for row in book[title].iter_rows(min_row=2,max_col=4):
                    if row[0].value==reference:nested[row[1].value]=(row[2].value,decoded_cell(row[3],metadata))
        assert nested['/nested/0']==('boolean',False)
        assert nested['/nested/1']==('number',0)
        assert nested['/nested/2']==('null',None)
        assert nested['/nested/3']==('number',9007199254740993)
        assert nested['/nested/4']==('number',0.12345678901234566)
        original=entity['observations'][-1]
        matched=False
        for title in book.sheetnames:
            if title.startswith('Historico'):
                columns={cell.value:cell.column-1 for cell in book[title][1]}
                for row in book[title].iter_rows(min_row=2):
                    if row[columns['id']].value==original['id']:
                        for field in ['value','input_value','previous']:
                            assert decoded_cell(row[columns[field]],metadata)==original[field]
                        matched=True
        assert matched
    finally:book.close()


@pytest.mark.parametrize('corruption',['converted_integer','boolean_as_zero','safe_number'])
def test_export_rejects_value_corruption_even_when_row_counts_match(tmp_path,monkeypatch,corruption):
    import bigbase.exports as exports_module
    exporter,job,_=make_numeric_export(tmp_path)
    actual_load=exports_module.load_workbook
    def corrupted_load(path,**kwargs):
        book=actual_load(path,read_only=False,data_only=False)
        sheet=book['Campos adicionais']
        columns={cell.value:cell.column for cell in sheet[1]}
        field={'converted_integer':'large_integer','boolean_as_zero':'false_value','safe_number':'safe_integer'}[corruption]
        cell=sheet.cell(2,columns[field])
        cell.value={'converted_integer':'9007199254740994','boolean_as_zero':0,'safe_number':12}[corruption]
        return book
    monkeypatch.setattr(exports_module,'load_workbook',corrupted_load)
    exporter.run(job['id'])
    with exporter.store.transaction() as tx:result=exporter.store.get(tx,'job',job['id'])
    assert result['status']=='failed'
    assert 'Conteúdo numérico' in result['error']
    assert not (exporter.root/(job['id']+'.xlsx')).exists()


def sheet_records(book,prefix):
    records=[]
    for title in book.sheetnames:
        if title==prefix or title.startswith(prefix+'_'):
            rows=book[title].iter_rows(values_only=True)
            headers=next(rows)
            records.extend(dict(zip(headers,row)) for row in rows)
    return records


def test_xlsx_preserves_catalog_versions_after_rename_deactivation_and_job_cutoff(tmp_path):
    from bigbase.catalogs import FieldDefinitionInput, create_definition, update_definition, validate_custom, bind_definition, VALUE_POLICY
    store=Store(tmp_path/'catalog.sqlite')
    field_id='synthetic-catalog-status'
    entity={'id':'synthetic-catalog-entity','entity_type':'person','version':2,'created_at':now(),'updated_at':now(),'items':[],'observations':[]}
    item={'id':'synthetic-catalog-item','kind':'custom','version':2,'fields':{},'sources':[],'notes':[]}
    entity['items'].append(item)
    def observe(tx,target,value,at,fields=None):
        metadata=validate_custom(store,tx,value,'person')
        start=len(entity['observations'])
        for key,data in value.items():
            if fields is None or key in fields:
                add_observation(entity,target,'value.'+key,data,'synthetic',at,'tester','catalog-op','',now())
        bind_definition(target,entity['observations'][start:],metadata)
    with store.transaction() as tx:
        create_definition(store,tx,FieldDefinitionInput(id=field_id,name='Nome Sintético Original',type='boolean'),'tester')
        observe(tx,item,{'field_id':field_id,'value':False},'2026-01-01T00:00:00Z')
        update_definition(store,tx,field_id,{'name':'Nome Sintético Revisado'},'1','tester')
        observe(tx,item,{'field_id':field_id,'value':False},'2026-02-01T00:00:00Z',{'value'})
        update_definition(store,tx,field_id,{'active':False},'2','tester')
        pending={'id':'synthetic-pending-item','kind':'custom','version':1,'fields':{},'sources':[],'notes':[]}
        entity['items'].append(pending)
        observe(tx,pending,{'field_id':'synthetic-unknown-field','value':0},'2026-01-01T00:00:00Z')
        project(entity)
        job=prepare_job({'all_records':True},[entity],'synthetic-owner')
        store.put(tx,'job',job)
        # The worker must not include a catalog edit made after this job's cutoff.
        update_definition(store,tx,field_id,{'name':'Nome Após o Corte','active':True},'3','tester')
    exporter=Exporter(store,tmp_path/'exports',row_limit=4)
    book=export_book(exporter,job)
    try:
        items={record['item_id']:record for record in sheet_records(book,'Metadados dos itens')}
        assert items[item['id']]['field_id']==field_id
        assert items[item['id']]['field_definition_version']==2
        assert items[item['id']]['classification_state']=='defined'
        assert items[item['id']]['value_policy']==VALUE_POLICY
        assert items[pending['id']]['classification_state']=='pending'
        assert items[pending['id']]['field_definition_version'] is None
        states={(record['item_id'],record['campo']):record for record in sheet_records(book,'Estado dos campos')}
        assert states[(item['id'],'value.field_id')]['field_definition_version']==1
        assert states[(item['id'],'value.value')]['field_definition_version']==2
        for original in entity['items']:
            for path,state in original['fields'].items():
                for key in ['field_id','field_definition_version','classification_state']:
                    assert states[(original['id'],path)][key]==state[key]
        observations={record['id']:record for record in sheet_records(book,'Historico')}
        for observation in entity['observations']:
            for key in ['field_id','field_definition_version','classification_state']:
                assert observations[observation['id']][key]==observation[key]
        definitions={row['field_definition_version']:row for row in sheet_records(book,'Definicoes dos campos')}
        assert set(definitions)=={1,2,3}
        assert definitions[1]['name']=='Nome Sintético Original'
        assert definitions[2]['name']=='Nome Sintético Revisado'
        assert definitions[3]['active'] is False
        assert all(row['field_id']==field_id for row in definitions.values())
    finally:book.close()
