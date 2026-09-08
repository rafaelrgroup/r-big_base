"""Synthetic acceptance tests for the local search/export ordering contract."""
from copy import deepcopy
import hashlib
import io
import json
import time
import zipfile

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from bigbase.exports import Exporter, prepare_job
from bigbase.sorting import resolve_sort, sort_entities
from test_flows import env, login, seed


def entity(id, name=None, addresses=(), birth=None, updated='2026-01-01T00:00:00Z'):
    return {'id': id, 'name': name, 'updated_at': updated, 'items': [
        {'kind': 'identity', 'value': {'name': name, 'birth_date': birth}, 'flags': {}},
        *[{'kind': 'address', 'value': value, 'flags': flags} for value, flags in addresses],
    ]}


def ordered(rows, sorts, **options):
    return [e['id'] for e in sort_entities(rows, resolve_sort({'sorts': sorts}), **options)]


def test_priorities_mixed_directions_and_stable_id():
    rows = [entity('b', 'Mesmo', [({'city': 'Alfa'}, {})], '1990-01-01'),
            entity('a', 'Mesmo', [({'city': 'Alfa'}, {})], '1990-01-01'),
            entity('d', 'Mesmo', [({'city': 'Beta'}, {})], '2005-01-01'),
            entity('c', 'Mesmo', [({'city': 'Alfa'}, {})], '2000-01-01')]
    original = deepcopy(rows)
    sorts = [{'field': 'city', 'direction': 'asc'}, {'field': 'birth_date', 'direction': 'desc'}]
    assert ordered(rows, sorts) == ['c', 'a', 'b', 'd']
    assert ordered(rows, list(reversed(sorts))) == ['d', 'c', 'a', 'b']
    assert rows == original


@pytest.mark.parametrize('direction,mode,expected', [
    ('asc', None, ['a', 'b']), ('desc', None, ['a', 'b']),
    ('asc', 'max', ['b', 'a']), ('desc', 'min', ['b', 'a']),
])
def test_multiple_addresses_default_and_explicit_modes(direction, mode, expected):
    rows = [entity('a', addresses=[({'city': 'Alfa'}, {}), ({'city': 'Zeta'}, {})]),
            entity('b', addresses=[({'city': 'Beta'}, {})])]
    criterion = {'field': 'city', 'direction': direction}
    if mode is not None: criterion['mode'] = mode
    assert ordered(rows, [criterion]) == expected


@pytest.mark.parametrize('field', ['name', 'birth_date', 'city', 'state', 'postal_code'])
@pytest.mark.parametrize('direction', ['asc', 'desc'])
def test_invalid_items_and_null_absent_values_last(field, direction):
    value = '2000-01-01' if field == 'birth_date' else '01234567'
    kind = 'identity' if field in {'name', 'birth_date'} else 'address'
    rows = [entity('c'), entity('b'), entity('a')]
    rows[0]['items'] = [{'kind': kind, 'value': {field: value}, 'flags': {'valid': False}}]
    rows[1]['items'] = [{'kind': kind, 'value': {field: None}, 'flags': {}}]
    rows[2]['items'] = [{'kind': kind, 'value': {field: value}, 'flags': {'valid': None}}]
    sorts = [{'field': field, 'direction': direction}]
    assert ordered(rows, sorts) == ['a', 'b', 'c']
    assert ordered(rows, sorts, include_invalid=True) == ['a', 'c', 'b']


def test_unicode_canonical_equivalence_and_empty_text():
    rows = [entity('b', 'Érica'), entity('a', 'E\u0301RICA'), entity('d'), entity('c', '')]
    assert ordered(rows, [{'field': 'name', 'direction': 'asc'}]) == ['c', 'a', 'b', 'd']
    assert ordered(rows, [{'field': 'name', 'direction': 'desc'}]) == ['a', 'b', 'c', 'd']


def test_postal_codes_are_text_and_wrong_json_types_are_missing():
    rows = [entity(str(i), addresses=[({'postal_code': v}, {})])
            for i, v in enumerate(['01234567', '10000000', '00999999', 'SW1A 1AA', 0, False])]
    assert ordered(rows, [{'field': 'postal_code', 'direction': 'asc'}]) == ['2', '0', '1', '3', '4', '5']


@pytest.mark.parametrize('field,values', [
    ('birth_date', ['2024-02-29', '2023-02-29', '01/01/2000', None, '20240101']),
    ('updated_at', ['2026-01-01T00:00:00Z', 'bad', '2026-01-01', None, '']),
])
@pytest.mark.parametrize('direction', ['asc', 'desc'])
def test_bad_dates_are_missing_without_discarding_record(field, values, direction):
    rows = [entity(str(i), birth=v, updated=v) for i, v in enumerate(values)]
    assert ordered(rows, [{'field': field, 'direction': direction}]) == ['0', '1', '2', '3', '4']


def test_timestamps_order_by_instant_and_explicit_id_direction():
    rows = [entity('b', updated='2026-01-01T00:00:00Z'), entity('a', updated='2025-12-31T21:00:00-03:00'),
            entity('c', updated='2025-12-31T22:00:00-03:00')]
    assert ordered(rows, [{'field': 'updated_at', 'direction': 'desc'}]) == ['c', 'a', 'b']
    assert ordered(rows, [{'field': 'updated_at', 'direction': 'asc'}, {'field': 'id', 'direction': 'desc'}]) == ['b', 'a', 'c']


BAD_SORTS = [None, [], {}, 'name', [None], [{'field': 'name'}], [{'direction': 'asc'}],
    [{'field': 'name', 'direction': 'asc', 'unknown': 1}],
    [{'field': 'name', 'direction': 'ASC'}], [{'field': 'name', 'direction': []}],
    [{'field': [], 'direction': 'asc'}], [{'field': 'custom:synthetic', 'direction': 'asc'}],
    [{'field': 'name', 'direction': 'asc', 'mode': None}],
    [{'field': 'name', 'direction': 'asc', 'mode': 'median'}],
    [{'field': 'id', 'direction': 'asc', 'mode': 'min'}],
    [{'field': 'updated_at', 'direction': 'asc', 'mode': 'max'}],
    [{'field': 'name', 'direction': 'asc'}, {'field': 'name', 'direction': 'desc'}],
    [{'field': 'id', 'direction': 'asc'}, {'field': 'name', 'direction': 'asc'}],
    [{'field': field, 'direction': 'asc'} for field in ['name', 'city', 'state', 'birth_date', 'postal_code', 'id']],
]


@pytest.mark.parametrize('sorts', BAD_SORTS)
def test_strict_sort_contract(sorts):
    with pytest.raises(HTTPException) as error: resolve_sort({'sorts': sorts})
    assert error.value.status_code == 422


@pytest.mark.parametrize('extra', [{'sort': 'name'}, {'direction': 'asc'}, {'sort': None}, {'include_invalid': 1}])
def test_reject_mixed_legacy_and_nonboolean(extra):
    with pytest.raises(HTTPException): resolve_sort({'sorts': [{'field': 'name', 'direction': 'asc'}], **extra})


def test_five_user_criteria_allow_implicit_sixth_id():
    fields = ['name', 'city', 'state', 'birth_date', 'postal_code']
    plan = resolve_sort({'sorts': [{'field': f, 'direction': 'asc'} for f in fields]})
    assert len(plan['criteria']) == 6
    assert plan['criteria'][-1] == {'field': 'id', 'direction': 'asc'}


def test_legacy_name_and_descending_id_ties_unchanged():
    rows = [entity('b', 'Mesmo'), entity('a', 'Mesmo'), entity('d', None), entity('c', 'Zulu')]
    for direction in ['asc', 'desc']:
        body = {'sort': 'name', 'direction': direction}
        expected = sorted(rows, key=lambda e: (str(e.get('name', '')).casefold(), e['id']), reverse=direction == 'desc')
        assert sort_entities(rows, resolve_sort(body)) == expected
    assert resolve_sort({})['contract'] == 'legacy'
    assert sort_entities(rows, resolve_sort({}, default_legacy=False)) == rows


def seed_ordering(c):
    rows = []
    for n, (city, birth) in enumerate([('Beta', '1990-01-01'), ('Alfa', '1980-01-01'), ('Alfa', '2000-01-01'), ('Alfa', '2000-01-01')]):
        r = seed(c, {'items': [
            {'kind': 'identity', 'value': {'name': 'Pessoa Sintética Ordenada', 'birth_date': birth}},
            {'kind': 'address', 'value': {'city': city, 'postal_code': '01234567'}},
        ]}, key='sort-seed-' + str(n))
        assert r.status_code == 200, r.text
        rows.append(r.json())
    return rows


def book_ids(data):
    book = load_workbook(io.BytesIO(data), read_only=True)
    try:
        ids = [row[0] for sheet in book if sheet.title == 'Cadastros' or sheet.title.startswith('Cadastros_') for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert 'Ordenacao' in book.sheetnames
        return ids
    finally: book.close()


@pytest.mark.parametrize('legacy', [False, True])
def test_saved_reopened_search_pages_and_xlsx_api_have_same_order(env, legacy):
    c, app, _ = env
    login(env)
    rows = seed_ordering(c)
    order = {'sort': 'name', 'direction': 'desc'} if legacy else {'sorts': [
        {'field': 'city', 'direction': 'asc'}, {'field': 'birth_date', 'direction': 'desc'}]}
    saved = c.post('/api/v1/saved-searches', json={'name': 'Ordenação sintética', **order}, headers={'Idempotency-Key': 'save-sort'})
    assert saved.status_code == 201, saved.text
    reopened = next(s for s in c.get('/api/v1/saved-searches').json()['items'] if s['id'] == saved.json()['id'])
    assert reopened['applied_sort']['contract'] == ('legacy' if legacy else 'multi')
    assert ('sorts' in reopened) is not legacy
    result = c.post('/api/v1/people/search', json=reopened).json()
    ids = [e['id'] for e in result['items']]
    expected = sorted((e['id'] for e in rows), reverse=True) if legacy else [*sorted(e['id'] for e in rows[2:]), rows[1]['id'], rows[0]['id']]
    assert ids == expected
    pages = [c.post('/api/v1/people/search', json={**reopened, 'limit': 2, 'offset': offset}).json() for offset in (0, 2)]
    assert [e['id'] for page in pages for e in page['items']] == ids
    assert all(page['applied_sort'] == result['applied_sort'] for page in pages)
    response = c.post('/api/v1/bulk-queries', json={**order, 'all_records': True}, headers={'Idempotency-Key': 'sort-export'})
    assert response.status_code == 202, response.text
    job = response.json()
    for _ in range(200):
        job = c.get('/api/v1/bulk-queries/' + job['id']).json()
        if job['status'] in {'completed', 'failed'}: break
        time.sleep(.02)
    assert job['status'] == 'completed', job
    assert job['applied_sort'] == result['applied_sort']
    downloaded = c.get('/api/v1/bulk-queries/' + job['id'] + '/files/result')
    assert downloaded.status_code == 200
    assert book_ids(downloaded.content) == ids


def test_contract_rejected_consistently_in_all_endpoints(env):
    c, _, _ = env
    login(env)
    for sorts in BAD_SORTS:
        for route in ['/people/search', '/companies/search', '/saved-searches', '/bulk-queries']:
            r = c.post('/api/v1' + route, json={'sorts': sorts, 'name': 'Teste', 'all_records': True}, headers={'Idempotency-Key': 'invalid-sort'})
            assert r.status_code == 422, (route, sorts, r.status_code, r.text)
    catalog = c.get('/api/v1/search/catalog').json()['sorting']
    assert {f['field'] for f in catalog['fields']} == {'name', 'id', 'updated_at', 'birth_date', 'city', 'state', 'postal_code'}
    assert all(f['status'] == 'ready_local' for f in catalog['fields'])


@pytest.mark.parametrize('include_invalid', [False, True])
def test_export_order_frozen_and_resumed_across_volumes(env, include_invalid):
    c, app, _ = env
    login(env)
    rows = seed_ordering(c)
    address = next(item for item in rows[0]['items'] if item['kind'] == 'address')
    r = c.patch(f"/api/v1/people/{rows[0]['id']}/items/{address['id']}", json={
        'source_id': 'manual', 'observed_at': '2026-02-01T00:00:00Z', 'flags': {'valid': False},
    }, headers={'If-Match': str(address['version'])})
    assert r.status_code == 200
    body = {'all_records': True, 'sorts': [{'field': 'city', 'direction': 'desc'}], 'include_invalid': include_invalid}
    expected = c.post('/api/v1/people/search', json=body).json()
    with app.state.store.transaction() as tx:
        job = prepare_job(body, app.state.store.all(tx, 'entity'), 'synthetic-owner')
        app.state.store.put(tx, 'job', job)
    body['sorts'][0]['direction'] = 'asc'
    assert job['criteria']['sorts'][0]['direction'] == 'desc'
    # A later write cannot alter the materialized selection or its values.
    r = seed(c, {'entity_id': rows[1]['id'], 'observed_at': '2026-02-01T00:00:00Z',
        'items': [{'kind': 'address', 'value': {'city': 'Zeta'}}]}, key='after-cutoff')
    assert r.status_code == 200
    exporter = Exporter(app.state.store, app.state.exports.root, row_limit=3, volume_limit=2)
    original = exporter.write_volume
    def crash(*args, **kwargs):
        original(*args, **kwargs)
        raise SystemExit('Synthetic worker loss after first volume')
    exporter.write_volume = crash
    with pytest.raises(SystemExit): exporter.run(job['id'])
    with app.state.store.transaction() as tx:
        interrupted = app.state.store.get(tx, 'job', job['id'])
        assert interrupted['status'] == 'preparing' and interrupted['snapshot']
    Exporter(app.state.store, app.state.exports.root, row_limit=3, volume_limit=2).run(job['id'])
    with app.state.store.transaction() as tx: final = app.state.store.get(tx, 'job', job['id'])
    assert final['status'] == 'completed', final
    assert final['cutoff'] == job['cutoff'] and final['snapshot'] == []
    ids = []
    with zipfile.ZipFile(app.state.exports.root / (job['id'] + '.zip')) as archive:
        manifest = json.loads(archive.read('manifesto.json'))
        assert manifest['applied_sort'] == expected['applied_sort']
        for volume in manifest['volumes']:
            data = archive.read(volume['file'])
            assert hashlib.sha256(data).hexdigest() == volume['sha256']
            ids.extend(book_ids(data))
    assert ids == [e['id'] for e in expected['items']]
    assert len(set(ids)) == len(rows)


def test_old_pending_export_without_order_preserves_snapshot_order(env):
    c, app, _ = env
    login(env)
    rows = seed_ordering(c)
    with app.state.store.transaction() as tx:
        job = prepare_job({'all_records': True}, list(reversed(rows)), 'synthetic-owner')
        job.pop('applied_sort')  # Old persistent jobs did not contain this property.
        app.state.store.put(tx, 'job', job)
    app.state.exports.run(job['id'])
    with app.state.store.transaction() as tx: final = app.state.store.get(tx, 'job', job['id'])
    assert final['status'] == 'completed', final
    assert final['applied_sort']['contract'] == 'snapshot_order'
    assert book_ids((app.state.exports.root / (job['id'] + '.xlsx')).read_bytes()) == [e['id'] for e in reversed(rows)]
