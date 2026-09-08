import io
import time
import json
import sqlite3
import pyotp
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from bigbase.api import create_app
from bigbase.domain import normalize, cpf_valid
from bigbase.search import matches, validate_filter

@pytest.fixture
def env(tmp_path):
    app=create_app(tmp_path,testing=True)
    with app.state.store.transaction() as c:
        u=app.state.security.create_user(c,'admin','synthetic-password-123','admin')
        secret=app.state.security.secret(u)
    with TestClient(app) as client:
        yield client,app,secret

def login(env):
    c,a,s=env
    r=c.post('/api/v1/auth/login',json={'username':'admin','password':'synthetic-password-123'});assert r.status_code==200,r.text
    j=r.json();r=c.post('/api/v1/auth/otp',json={'challenge':j['challenge'],'code':pyotp.TOTP(s).now()});assert r.status_code==200,r.text
    c.headers['X-CSRF-Token']=r.json()['csrf'];return r.json()

def seed(c,extra=None,key='seed'):
    body={'source_id':'manual','observed_at':'2026-01-01T00:00:00Z','items':[{'kind':'identity','value':{'name':'Pessoa Sintética Alfa','birth_date':'1990-01-01'}},{'kind':'phone','value':{'number':'+5511998765432'},'flags':{'is_whatsapp':None}}]}
    if extra:body.update(extra)
    return c.post('/api/v1/people/enrich',json=body,headers={'Idempotency-Key':key})

def test_otp_enrollment_and_replay(env):
    c,a,s=env
    r=c.post('/api/v1/auth/login',json={'username':'admin','password':'synthetic-password-123'})
    assert r.json()['enrollment_required'];assert c.post('/api/v1/people/search',json={}).status_code==401
    token=r.json()['challenge'];r=c.post('/api/v1/auth/otp',json={'challenge':token,'code':pyotp.TOTP(s).now()});assert r.status_code==200
    assert len(r.json()['recovery_codes'])==10
    assert c.post('/api/v1/auth/otp',json={'challenge':token,'code':pyotp.TOTP(s).now()}).status_code==401

def test_csrf(env):
    login(env);c=env[0];c.headers.pop('X-CSRF-Token')
    assert seed(c).status_code==403

def test_idempotency_and_history(env):
    login(env);c=env[0];a=seed(c);assert a.status_code==200,a.text
    b=seed(c);assert b.json()==a.json()
    assert seed(c,{'reason':'different'}).status_code==409
    assert len(c.post('/api/v1/people/search',json={}).json()['items'])==1
    hist=c.get('/api/v1/people/'+a.json()['id']+'/history').json()['items'];assert hist and all(o['source_id']=='manual' for o in hist)

def test_late_observation_and_partial_enrichment(env):
    login(env);c=env[0];e=seed(c).json()
    r=seed(c,{'entity_id':e['id'],'observed_at':'2025-01-01T00:00:00Z','items':[{'kind':'identity','value':{'name':'Nome Antigo'}}]},key='late')
    assert r.status_code==200,r.text
    assert r.json()['name']=='Pessoa Sintética Alfa'
    assert any(not o['applied'] and o['value']=='Nome Antigo' for o in r.json()['observations'])
    assert any(x['kind']=='phone' for x in r.json()['items'])

def test_unknown_date_never_overrides_dated(env):
    login(env);c=env[0];e=seed(c).json()
    r=seed(c,{'entity_id':e['id'],'observed_at':None,'items':[{'kind':'identity','value':{'name':'Sem Data'}}]},key='undated')
    assert r.json()['name']==e['name']

def test_patch_flags_versions_and_history(env):
    login(env);c=env[0];e=seed(c).json();it=next(x for x in e['items'] if x['kind']=='phone');url=f"/api/v1/people/{e['id']}/items/{it['id']}"
    body={'source_id':'manual','observed_at':'2026-02-01T00:00:00Z','flags':{'is_whatsapp':False,'valid':False}}
    assert c.patch(url,json=body).status_code==428
    r=c.patch(url,json=body,headers={'If-Match':str(it['version'])});assert r.status_code==200,r.text
    assert next(x for x in r.json()['items'] if x['id']==it['id'])['flags']['is_whatsapp'] is False
    assert c.patch(url,json=body,headers={'If-Match':str(it['version'])}).status_code==409
    assert len(r.json()['observations'])>len(e['observations'])

def test_same_item_filter(env):
    login(env);c=env[0]
    seed(c,{'items':[{'kind':'identity','value':{'name':'Pessoa Sintética'}},{'kind':'address','value':{'city':'Cidade Alfa','postal_code':'11111000'}},{'kind':'address','value':{'city':'Cidade Beta','postal_code':'22222000'}}]})
    filters={'item':{'and':[{'field':'city','value':'Cidade Alfa'},{'field':'postal_code','value':'22222000'}]}}
    assert c.post('/api/v1/people/search',json={'filters':filters}).json()['total']==0
    filters['item']['and'][1]['value']='11111000'
    assert c.post('/api/v1/people/search',json={'filters':filters}).json()['total']==1

def test_combined_name_search(env):
    login(env);c=env[0];seed(c)
    r=c.post('/api/v1/people/search',json={'filters':{'and':[{'field':'name','op':'contains','value':'sintetica'},{'field':'entity_type','value':'person'}]}})
    assert r.status_code==200,r.text
    assert r.json()['total']==1

def test_invalid_filter_rejected(env):
    login(env);c=env[0]
    assert c.post('/api/v1/people/search',json={'filters':{'field':'made_up','value':'x'}}).status_code==422

@pytest.mark.parametrize('v',[True,False,None])
def test_flags_strict_and_tristate(env,v):
    login(env);c=env[0]
    r=seed(c,{'items':[{'kind':'phone','value':{'number':'+5511998765432'},'flags':{'is_whatsapp':v}}]})
    assert r.status_code==200;assert r.json()['items'][0]['flags']['is_whatsapp'] is v

@pytest.mark.parametrize('v',['false',0,1,'yes'])
def test_flags_no_coercion(env,v):
    login(env);c=env[0]
    assert seed(c,{'items':[{'kind':'phone','value':{'number':'123'},'flags':{'valid':v}}]}).status_code==422

def test_xlsx_full_export(env):
    login(env);c=env[0]
    e=seed(c,{'items':[{'kind':'identity','value':{'name':'=HYPERLINK("x")'}},{'kind':'phone','value':{'number':'+5511998765432'},'flags':{'is_whatsapp':False}},{'kind':'phone','value':{'number':'0123'},'flags':{'is_whatsapp':None}}]}).json()
    r=c.post('/api/v1/bulk-queries',json={'all_records':True},headers={'Idempotency-Key':'bulk'});assert r.status_code==202,r.text
    id=r.json()['id']
    for _ in range(100):
        j=c.get('/api/v1/bulk-queries/'+id).json()
        if j['status'] in {'completed','failed'}:break
        time.sleep(.02)
    assert j['status']=='completed',j
    r=c.get(f'/api/v1/bulk-queries/{id}/files/result');assert r.status_code==200,r.text
    wb=load_workbook(io.BytesIO(r.content));assert wb['Cadastros']['C2'].value=='=HYPERLINK("x")';assert wb['Cadastros']['C2'].data_type=='s'
    assert wb['Telefones'].max_row==3;assert wb['Historico'].max_row==len(e['observations'])+1

def test_append_only_audit(env):
    login(env);_,a,_=env
    with pytest.raises(sqlite3.IntegrityError):
        with a.state.store.transaction() as c:c.execute('DELETE FROM events')

@pytest.mark.parametrize('kind,value,key,expected',[
 ('document',{'type':'CPF','number':'000.000.000-00'},'number','00000000000'),
 ('email',{'email':'User+Tag@EXAMPLE.COM'},'email','User+Tag@example.com'),
 ('phone',{'number':'+551133334444'},'classification','fixed'),
 ('phone',{'number':'+5511998765432'},'classification','mobile'),
 ('phone',{'number':'1234'},'number','1234'),
 ('address',{'postal_code':'01234-567'},'postal_code','01234567')])
def test_normalization(kind,value,key,expected):
    assert normalize(kind,value)[0][key]==expected

def test_recovery_revokes_session_and_requires_new_enrollment(env):
    data=login(env);c,a,s=env
    challenge=c.post('/api/v1/auth/login',json={'username':'admin','password':'synthetic-password-123'}).json()['challenge']
    r=c.post('/api/v1/auth/recovery',json={'challenge':challenge,'code':data['recovery_codes'][0]})
    assert r.status_code==200,r.text
    assert r.json()['enrollment_required'];assert 'otp_uri' in r.json()
    assert c.get('/api/v1/auth/me').status_code==401
    assert c.post('/api/v1/auth/recovery',json={'challenge':challenge,'code':data['recovery_codes'][0]}).status_code==401

def test_source_date_has_precedence_and_preserved(env):
    login(env);c=env[0];e=seed(c).json()
    r=seed(c,{'entity_id':e['id'],'observed_at':'2025-01-01T00:00:00Z','source_updated_at':'2026-02-01T00:00:00Z','items':[{'kind':'identity','value':{'name':'Atualização Sintética'}}]},key='source-date')
    assert r.status_code==200,r.text
    assert r.json()['name']=='Atualização Sintética'
    obs=r.json()['observations'][-1]
    assert obs['source_observed_at']=='2025-01-01T00:00:00Z'
    assert obs['source_updated_at']=='2026-02-01T00:00:00Z'

def test_external_source_identity_stays_stable(env):
    login(env);c=env[0]
    e=seed(c,{'external_id':'external-synthetic-1'}).json()
    r=seed(c,{'external_id':'external-synthetic-1','items':[{'kind':'email','value':{'email':'new@example.invalid'}}]},key='source-v2')
    assert r.status_code==200,r.text
    assert r.json()['id']==e['id']

def test_api_key_scope_source_revocation(env):
    login(env);c,a,s=env
    with a.state.store.transaction() as tx:
        u=a.state.store.all(tx,'user')[0];u['last_otp_step']=-1;a.state.store.put(tx,'user',u)
    r=c.post('/api/v1/admin/api-keys',json={'name':'Synthetic key','scopes':['read'],'sources':['manual'],'otp':pyotp.TOTP(s).now()})
    assert r.status_code==200,r.text
    key=r.json()
    with TestClient(a) as other:
        headers={'X-API-Key':key['key'],'Idempotency-Key':'forbidden'}
        assert other.post('/api/v1/people/search',json={},headers=headers).status_code==200
        assert other.post('/api/v1/people/enrich',json={'source_id':'manual','items':[{'kind':'identity','value':{'name':'Forbidden'}}]},headers=headers).status_code==403
        assert c.patch('/api/v1/admin/api-keys/'+key['id']).status_code==200
        assert other.post('/api/v1/people/search',json={},headers=headers).status_code==401

def test_future_evidence_is_pending(env):
    login(env);c=env[0]
    e=seed(c,{'observed_at':'2099-01-01T00:00:00Z'}).json()
    assert all(o['pending_reason']=='future_timestamp' for o in e['observations'])
    assert e['name']=='Sem nome informado'

def test_missing_source_rejected(env):
    login(env);c=env[0]
    assert seed(c,{'source_id':'unregistered'}).status_code==422

def test_wrong_patch_date_is_validation_error(env):
    login(env);c=env[0];e=seed(c).json();it=e['items'][0]
    r=c.patch(f"/api/v1/people/{e['id']}/items/{it['id']}",json={'source_id':'manual','observed_at':'wrong','flags':{'valid':True}},headers={'If-Match':str(it['version'])})
    assert r.status_code==422

def test_upload_csv_and_bulk_lookup(env):
    login(env);c=env[0];seed(c)
    r=c.post('/api/v1/bulk-queries/uploads',files={'file':('names.csv','Pessoa Sintética Alfa\nNão Encontrado\n'.encode(),'text/csv')})
    assert r.status_code==201,r.text
    assert r.json()['count']==2
    j=c.post('/api/v1/bulk-queries',json={'upload_id':r.json()['id'],'input_field':'name'},headers={'Idempotency-Key':'upload-job'})
    assert j.status_code==202,j.text
    for _ in range(100):
        status=c.get('/api/v1/bulk-queries/'+j.json()['id']).json()
        if status['status'] in {'completed','failed'}:break
        time.sleep(.02)
    assert status['status']=='completed',status
    assert status['entity_count']==1

def test_age_feb29_and_casefold():
    from datetime import date
    entity={'items':[{'kind':'identity','value':{'birth_date':'2000-02-29'},'flags':{},'sources':[]}],'name':'Álfa Sintética'}
    assert matches(entity,{'field':'age','value':24},reference_date=date(2025,2,28))
    assert matches(entity,{'field':'age','value':25},reference_date=date(2025,3,1))
    assert matches(entity,{'field':'name','value':'alfa sintetica'})

def test_xlsx_structures_and_sheet_split(env):
    login(env);c,a,_=env
    seed(c,{'items':[{'kind':'custom','value':{'field_id':'complex','value':{'nested':[False,None,'000123']}}}]})
    with a.state.store.transaction() as tx:
        from bigbase.exports import prepare_job,Exporter
        u=a.state.store.all(tx,'user')[0]
        j=prepare_job({'all_records':True},a.state.store.all(tx,'entity'),u['id']);a.state.store.put(tx,'job',j)
    exporter=Exporter(a.state.store,a.state.exports.root,row_limit=4);exporter.run(j['id'])
    with a.state.store.transaction() as tx:result=a.state.store.get(tx,'job',j['id'])
    assert result['status']=='completed',result.get('error')
    assert result['snapshot']==[]
    assert any(name.startswith('Valores estruturados_') for name in result['sheet_counts'])
    wb=load_workbook(a.state.exports.root/(j['id']+'.xlsx'),read_only=True)
    values=[row[3] for name in wb.sheetnames if name.startswith('Valores estruturados') for row in wb[name].iter_rows(min_row=2,max_col=4,values_only=True)]
    assert '000123' in values and False in values and None in values


def test_official_alphanumeric_cnpj_vector():
    # Receita Federal, Manual de cálculo do DV do CNPJ, exemplo 12.ABC.345/01DE-35.
    assert normalize('document',{'type':'CNPJ','number':'12.ABC.345/01DE-35'})[0]['syntax_valid'] is True
    assert normalize('document',{'type':'CNPJ','number':'12.ABC.345/01DE-34'})[0]['syntax_valid'] is False

def test_future_document_does_not_merge_identity(env):
    login(env);c=env[0]
    doc={'kind':'document','value':{'type':'CNPJ','number':'12.ABC.345/01DE-35'}}
    original=seed(c,{'items':[doc]}).json()
    future=seed(c,{'items':[doc],'observed_at':'2099-01-01T00:00:00Z'},key='future-doc').json()
    assert future['id'] != original['id']
    assert all(o['pending_reason']=='future_timestamp' for o in future['observations'])

def test_malformed_flags_and_typed_filter_lists(env):
    login(env);c=env[0]
    assert seed(c,{'items':[{'kind':'phone','value':{'number':'123'},'flags':[]}]}).status_code==422
    for f in [{'field':'is_whatsapp','op':'in','value':[1]}, {'field':'email','op':'exists','value':'false'}]:
        assert c.post('/api/v1/people/search',json={'filters':f}).status_code==422


def test_confirmation_expiry_preserves_known_result_and_independent_flags(env):
    login(env);c=env[0]
    e=seed(c,{'items':[{'kind':'phone','value':{'number':'+5511998765432'},'flags':{'valid':True},'flag_evidence':{'is_whatsapp':{'value':False,'checked_at':'2025-01-01T00:00:00Z','expires_at':'2025-02-01T00:00:00Z','method':'synthetic-test'}}}]}).json()
    it=e['items'][0]
    assert it['flags']=={'valid':True,'is_whatsapp':False}
    assert it['flag_details']['is_whatsapp']['stale'] is True
    assert it['flag_details']['valid']['stale'] is False
    r=c.patch(f"/api/v1/people/{e['id']}/items/{it['id']}",json={'source_id':'manual','flag_evidence':{'is_whatsapp':{'value':True,'checked_at':'2026-01-02T00:00:00Z','method':'synthetic-recheck'}}},headers={'If-Match':str(it['version'])})
    assert r.status_code==200,r.text
    updated=r.json()['items'][0]
    assert updated['flag_details']['valid']==it['flag_details']['valid']
    assert updated['flag_details']['is_whatsapp']['stale'] is False
    assert any(o.get('verification',{}).get('expires_at')=='2025-02-01T00:00:00Z' for o in r.json()['observations'])

@pytest.mark.parametrize('evidence',[
 {'value':True,'expires_at':'2026-01-02T00:00:00Z'},
 {'value':None,'checked_at':'2026-01-01T00:00:00Z','expires_at':'2026-01-02T00:00:00Z'},
 {'value':True,'checked_at':'2026-01-01T00:00:00Z','expires_at':'2025-01-02T00:00:00Z'},
 {'value':'false'},
])
def test_confirmation_invalid_contract_rejected(env,evidence):
    login(env);c=env[0]
    assert seed(c,{'items':[{'kind':'phone','value':{'number':'123'},'flag_evidence':{'is_whatsapp':evidence}}]}).status_code==422

def test_explicit_null_input_is_preserved(env):
    login(env);c=env[0]
    e=seed(c,{'items':[{'kind':'document','value':{'type':'OTHER','number':None}}]}).json()
    obs=next(o for o in e['observations'] if o['path']=='value.number')
    assert obs['input_value'] is None


def test_xlsx_multiple_volumes_columns_and_manifest(env):
    import zipfile,hashlib
    from bigbase.exports import Exporter,prepare_job
    login(env);c,a,_=env
    value={('=UNTRUSTED_HEADER()' if n==0 else 'column_%02d'%n):'value_%02d'%n for n in range(42)}
    first=seed(c,{'items':[{'kind':'identity','value':{'name':'Pessoa Volume Um',**value}}]}).json()
    second=seed(c,{'items':[{'kind':'identity','value':{'name':'Pessoa Volume Dois'}}]},key='second-volume').json()
    with a.state.store.transaction() as tx:
        u=a.state.store.all(tx,'user')[0]
        j=prepare_job({'all_records':True},a.state.store.all(tx,'entity'),u['id']);a.state.store.put(tx,'job',j)
    exporter=Exporter(a.state.store,a.state.exports.root,volume_limit=1,column_limit=32);exporter.run(j['id'])
    with a.state.store.transaction() as tx:result=a.state.store.get(tx,'job',j['id'])
    assert result['status']=='completed',result.get('error')
    assert result['artifact']=='zip' and len(result['volumes'])==2
    response=c.get('/api/v1/bulk-queries/'+j['id']+'/files/result')
    assert response.status_code==200 and response.headers['content-type']=='application/zip'
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        manifest=json.loads(archive.read('manifesto.json'))
        ids=set();preserved={}
        for volume in manifest['volumes']:
            data=archive.read(volume['file']);assert hashlib.sha256(data).hexdigest()==volume['sha256']
            book=load_workbook(io.BytesIO(data))
            ids.add(book['Cadastros']['A2'].value)
            for name in book.sheetnames:
                if name.startswith('Identidade'):
                    assert book[name].max_column<=32
                    headers=list(book[name][1]);values=list(book[name][2])
                    for h,v in zip(headers,values):
                        if h.value in value:preserved[h.value]=v.value;assert h.data_type=='s'
        assert ids=={first['id'],second['id']}
        assert preserved==value

def test_bulk_duplicate_input_ids_rejected(env):
    login(env);c=env[0]
    r=c.post('/api/v1/bulk-queries',json={'entries':[{'input_id':'same','value':'First'},{'input_id':'same','value':'Second'}]},headers={'Idempotency-Key':'duplicate-input'})
    assert r.status_code==422


def test_relationships_inverse_and_pending_preserve_original_event(env):
    login(env);c=env[0]
    mother=seed(c,{'items':[{'kind':'identity','value':{'name':'Pessoa Sintética Mãe'}}]}).json()
    child=seed(c,{'items':[{'kind':'identity','value':{'name':'Pessoa Sintética Filha'}},{'kind':'relationship','value':{'target_id':mother['id'],'type':'mother'},'flags':{'valid':None}},{'kind':'relationship','value':{'target_name':'Possível Parente Sintético','type':'possible_relative'}}]},key='child').json()
    outgoing=c.get('/api/v1/people/'+child['id']+'/relationships').json()['outgoing']
    assert {x['resolution']['status'] for x in outgoing}=={'resolved','pending_identity'}
    incoming=c.get('/api/v1/people/'+mother['id']+'/relationships').json()['incoming']
    assert len(incoming)==1 and incoming[0]['inverse_role']=='child'
    assert incoming[0]['derived'] is True
    assert incoming[0]['item']['flags']['valid'] is None
    assert incoming[0]['observations'][0]['source_id']=='manual'
    assert len(c.get('/api/v1/people/'+mother['id']).json()['observations'])==len(mother['observations'])
    assert seed(c,{'entity_id':mother['id'],'items':[{'kind':'relationship','value':{'target_id':mother['id'],'type':'mother'}}]},key='self-link').status_code==422


def test_invitation_one_use_and_otp_required(env):
    login(env);c,a,_=env
    r=c.post('/api/v1/admin/invitations',json={'username':'invited-user','permissions':['read']})
    assert r.status_code==201,r.text
    token=r.json()['activation_token']
    with TestClient(a) as user:
        assert user.post('/api/v1/auth/login',json={'username':'invited-user','password':'anything'}).status_code==401
        result=user.post('/api/v1/auth/activate',json={'token':token,'password':'invited-synthetic-123'})
        assert result.status_code==200 and result.json()['enrollment_required']
        assert user.post('/api/v1/people/search',json={}).status_code==401
        assert user.post('/api/v1/auth/activate',json={'token':token,'password':'another-password-123'}).status_code==401
        with a.state.store.transaction() as tx:
            u=next(u for u in a.state.store.all(tx,'user') if u['username']=='invited-user');secret=a.state.security.secret(u)
        result=user.post('/api/v1/auth/otp',json={'challenge':result.json()['challenge'],'code':pyotp.TOTP(secret).now()})
        assert result.status_code==200
        user.headers['X-CSRF-Token']=result.json()['csrf']
        assert user.post('/api/v1/people/search',json={}).status_code==200
        assert user.post('/api/v1/admin/invitations',json={'username':'forbidden'}).status_code==403
        assert c.patch('/api/v1/admin/users/'+u['id'],json={'active':False}).status_code==200
        assert c.patch('/api/v1/admin/users/'+u['id'],json={'active':True}).status_code==200
        assert user.post('/api/v1/people/search',json={}).status_code==401


def test_saved_search_ownership_idempotency_and_archiving(env):
    login(env);c,a,_=env
    filters={'item':{'or':[{'field':'city','value':'Cidade A'},{'field':'city','value':'Cidade B'}]}}
    body={'name':'Cidades sintéticas','filters':filters,'include_invalid':True}
    r=c.post('/api/v1/saved-searches',json=body,headers={'Idempotency-Key':'saved'})
    assert r.status_code==201,r.text
    assert c.post('/api/v1/saved-searches',json=body,headers={'Idempotency-Key':'saved'}).json()['id']==r.json()['id']
    assert c.get('/api/v1/saved-searches').json()['items'][0]['filters']==filters
    assert c.patch('/api/v1/saved-searches/'+r.json()['id']).status_code==200
    assert c.get('/api/v1/saved-searches').json()['items']==[]
    with a.state.store.transaction() as tx:assert a.state.store.get(tx,'saved_search',r.json()['id'])['filters']==filters


@pytest.mark.parametrize('body',[
 {'username':'bad','password':None},
 {'username':123,'password':'long-enough-password'},
 {'username':'bad','password':'long-enough-password','permissions':[['read']]},
])
def test_malformed_user_input_returns_422(env,body):
    login(env);c=env[0]
    assert c.post('/api/v1/admin/users',json=body).status_code==422

def test_nonfinite_json_rejected_without_partial_write(env):
    login(env);c=env[0]
    assert c.post('/api/v1/people/enrich',content='{"source_id":"manual","items":[{"kind":"identity","value":{"weight":NaN}}]}',headers={'Content-Type':'application/json','Idempotency-Key':'nan'}).status_code==422
    assert c.post('/api/v1/people/search',json={}).json()['total']==0


def test_confirmation_does_not_transfer_to_changed_identity_value(env):
    login(env);c=env[0]
    e=seed(c,{'items':[{'kind':'identity','value':{'name':'Primeiro Nome Sintético'},'flags':{'valid':True}}]}).json()
    r=seed(c,{'entity_id':e['id'],'observed_at':'2026-02-01T00:00:00Z','items':[{'kind':'identity','value':{'name':'Segundo Nome Sintético'}}]},key='change-identity')
    assert r.status_code==200,r.text
    it=r.json()['items'][0]
    assert it['flags']['valid'] is None
    assert it['flag_details']['valid']['last_result'] is True
    assert it['flag_details']['valid']['applicable'] is False
    assert any(o['path']=='flag.valid' and o['value'] is True for o in r.json()['observations'])


def test_normalized_document_phone_search_and_cnpj_bulk(env):
    from bigbase.exports import parse_entries
    login(env);c=env[0]
    seed(c,{'items':[{'kind':'document','value':{'type':'CNPJ','number':'12.ABC.345/01DE-35'}},{'kind':'phone','value':{'number':'+5511998765432'}}]})
    for field,value in [('document','12.ABC.345/01DE-35'),('phone','(11) 99876-5432')]:
        r=c.post('/api/v1/people/search',json={'filters':{'field':field,'value':value}})
        assert r.status_code==200 and r.json()['total']==1,r.text
    entries=parse_entries({'text':'12.ABC.345/01DE-35','input_field':'document','document_type':'CNPJ'})
    assert entries[0]['error'] is None
    r=c.post('/api/v1/people/search',json={'filters':entries[0]['filter']})
    assert r.json()['total']==1


def test_expired_export_removes_only_temporary_file(env):
    from bigbase.exports import prepare_job
    login(env);c,a,_=env;entity=seed(c).json()
    with a.state.store.transaction() as tx:
        u=a.state.store.all(tx,'user')[0]
        job=prepare_job({'all_records':True},a.state.store.all(tx,'entity'),u['id']);a.state.store.put(tx,'job',job)
    a.state.exports.run(job['id'])
    path=a.state.exports.root/(job['id']+'.xlsx');assert path.exists()
    with a.state.store.transaction() as tx:
        j=a.state.store.get(tx,'job',job['id']);j['expires_at']='2000-01-01T00:00:00Z';a.state.store.put(tx,'job',j)
    assert a.state.exports.cleanup_expired()==1
    assert not path.exists()
    assert c.get('/api/v1/bulk-queries/'+job['id']+'/files/result').status_code==410
    assert c.get('/api/v1/people/'+entity['id']).json()['observations']==entity['observations']
    assert c.get('/api/v1/bulk-queries/'+job['id']).json()['artifact_expired'] is True

def test_collection_mismatch_does_not_mutate_other_type(env):
    login(env);c=env[0];entity=seed(c).json();item=entity['items'][0]
    assert c.get('/api/v1/companies/'+entity['id']).status_code==404
    assert c.get('/api/v1/companies/'+entity['id']+'/history').status_code==404
    assert c.patch('/api/v1/companies/'+entity['id']+'/items/'+item['id'],json={'source_id':'manual','flags':{'valid':True}},headers={'If-Match':str(item['version'])}).status_code==404


def test_normalizer_version_and_original_format_are_auditable(env):
    login(env);c=env[0]
    e=seed(c,{'items':[{'kind':'phone','value':{'number':'(11) 99876-5432'}}]}).json()
    obs=next(o for o in e['observations'] if o['path']=='value.number')
    assert obs['input_value']=='(11) 99876-5432'
    assert obs['value']=='+5511998765432'
    assert obs['normalization']['changed'] is True
    assert obs['normalization']['input_path']=='/items/0/value/number'
    assert 'phonenumbers-' in obs['normalization']['version']
    assert next(o for o in e['observations'] if o['path']=='value.classification')['normalization']['derived'] is True
